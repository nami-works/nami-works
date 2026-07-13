"""Build the Beleza na Web SKU registry .xlsx (Outcome 1) + _Pendencias.md (Outcome 3).

Reads _map.json (mapped + unmapped SKUs) and _labels.json (final curated labels).
Writes both deliverables to the GEB shared Drive root. House style per
docs/excel-conventions.md: Calibri 8pt, red #DF3630 headers, no merged cells,
freeze panes, gridlines off, pt-BR labels.
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent / "imagery" / "blz-catalog"
MAP = json.load(open(BASE / "_map.json", encoding="utf-8"))
LABELS = json.load(open(BASE / "_labels.json", encoding="utf-8"))
DRIVE_ROOT = Path(
    r"G:\Drives compartilhados\GEB_Comercial\Marketplaces"
    r"\Beleza na Web\Cadastro GE Beauty BLZ\Imagens"
)
PHOTO_TYPES = ["still", "ambientada", "textura", "infografico"]

RED = PatternFill("solid", fgColor="DF3630")
GREY = PatternFill("solid", fgColor="F7F7F7")
AMBER = PatternFill("solid", fgColor="FFF2CC")
WHITE_BOLD = Font(name="Calibri", size=8, bold=True, color="FFFFFF")
BOLD = Font(name="Calibri", size=8, bold=True)
NORM = Font(name="Calibri", size=8)
THIN = Border(*[Side(style="thin", color="D9D9D9")] * 4)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)


def slugify(s):
    s = (s or "").lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:40]


def counts_for(sku):
    c = {t: 0 for t in PHOTO_TYPES}
    for e in LABELS.get(sku, []):
        c[e["label"]] = c.get(e["label"], 0) + 1
    return c


def build_registry(ws):
    headers = ["SKU", "EAN", "Produto", "Tipo", "Handle Shopify", "Pasta (Drive)",
               "Still", "Ambientada", "Textura", "Infográfico", "Vídeo",
               "620x620", "3000x3000", "Pendências"]
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h)
        c.fill = RED
        c.font = WHITE_BOLD
        c.alignment = CENTER
        c.border = THIN
    r = 2
    tot = {t: 0 for t in PHOTO_TYPES}
    for rec in MAP["mapped"]:
        sku = rec["sku"]
        c = counts_for(sku)
        for t in PHOTO_TYPES:
            tot[t] += c[t]
        folder = f"{rec['ean'].strip()}_{slugify(rec['name'])}"
        n620 = sum(c[t] for t in PHOTO_TYPES)
        pend = []
        if c["textura"] == 0:
            pend.append("sem textura")
        pend.append("vídeo")
        pend.append("upscale 3000")
        row = [sku, rec["ean"].strip(), rec["name"], rec.get("type", ""), rec["handle"],
               folder, c["still"], c["ambientada"], c["textura"], c["infografico"],
               "Pendente", f"OK ({n620})", "Pendente (upscale)", "; ".join(pend)]
        for j, v in enumerate(row, 1):
            cell = ws.cell(r, j, v)
            cell.font = NORM
            cell.border = THIN
            cell.fill = GREY
            cell.alignment = CENTER if 7 <= j <= 13 else LEFT
        r += 1
    # totals row
    ws.cell(r, 1, "TOTAL").font = BOLD
    ws.cell(r, 6, f"{len(MAP['mapped'])} SKUs").font = BOLD
    for off, t in zip(range(7, 11), PHOTO_TYPES):
        cc = ws.cell(r, off, tot[t]); cc.font = BOLD; cc.alignment = CENTER
    for j in range(1, 15):
        ws.cell(r, j).fill = AMBER
        ws.cell(r, j).border = THIN
    widths = [9, 15, 34, 10, 32, 40, 7, 11, 8, 12, 9, 10, 16, 26]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:N{r-1}"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 115
    ws.sheet_format.defaultRowHeight = 10.5
    ws.row_dimensions[1].height = 26


def build_pending_sheet(ws):
    hdr = ["SKU", "EAN", "Produto", "Motivo"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(1, j, h); c.fill = RED; c.font = WHITE_BOLD; c.alignment = CENTER; c.border = THIN
    r = 2
    for u in MAP["unmapped"]:
        row = [u["sku"], u.get("ean") or "—", u["name"], u["reason"]]
        for j, v in enumerate(row, 1):
            cell = ws.cell(r, j, v); cell.font = NORM; cell.border = THIN; cell.fill = GREY
            cell.alignment = LEFT
        r += 1
    for j, w in enumerate([9, 15, 40, 34], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.sheet_format.defaultRowHeight = 10.5
    ws.row_dimensions[1].height = 26


def build_md():
    L = []
    L.append("# GE Beauty — Beleza na Web | Itens pendentes\n")
    L.append("Gerado a partir do Shopify (fonte canônica). Este documento lista o que a "
             "etapa de download **não** conseguiu entregar.\n")
    L.append("## 1. Formato 3000x3000 — pendente para 100% dos SKUs")
    L.append("Nenhuma imagem no Shopify atinge 3000px (máximo observado ~2500px). "
             "O formato 620x620 foi entregue para todos. Em cada pasta `.../3000/` "
             "ficou salvo o **arquivo nativo de maior resolução disponível** "
             "(`*_NATIVE_WxH.jpg`) + marcador `_PENDENTE_upscale_3000.txt`, pronto para "
             "upscale (Magnific ou equivalente) quando aprovado.\n")
    L.append("## 2. Vídeos — pendente para 100% dos SKUs")
    L.append("Não há **nenhum** vídeo (textura/aplicação) no Shopify. Precisam ser "
             "produzidos ou fornecidos por outra fonte.\n")
    L.append("## 3. Fotos de textura — quase inexistentes")
    L.append("O Shopify praticamente não tem close-ups de textura isolados; a textura "
             "aparece apenas dentro de infográficos. SKUs sem nenhuma textura recebem "
             "marcador `_PENDENTE.txt` na pasta `textura/`.\n")
    zero_tex = [rec["sku"] for rec in MAP["mapped"] if counts_for(rec["sku"])["textura"] == 0]
    L.append(f"SKUs sem foto de textura ({len(zero_tex)}): " + ", ".join(zero_tex) + "\n")
    L.append("## 4. SKUs sem imagens (não mapeados no Shopify)")
    L.append("Estes SKUs do `products.json` não têm correspondência com imagens no "
             "Shopify e precisam de fotografia/cadastro à parte:\n")
    L.append("| SKU | EAN | Produto | Motivo |")
    L.append("|---|---|---|---|")
    for u in MAP["unmapped"]:
        L.append(f"| {u['sku']} | {u.get('ean') or '—'} | {u['name']} | {u['reason']} |")
    L.append("")
    return "\n".join(L)


def main():
    wb = Workbook()
    build_registry(wb.active)
    wb.active.title = "SKUs BLZ"
    build_pending_sheet(wb.create_sheet("Sem imagem"))
    out_xlsx = DRIVE_ROOT / "_Planilha_GE_Beauty_SKUs_BLZ.xlsx"
    wb.save(out_xlsx)
    (DRIVE_ROOT / "_Pendencias.md").write_text(build_md(), encoding="utf-8")
    print("wrote", out_xlsx)
    print("wrote", DRIVE_ROOT / "_Pendencias.md")


if __name__ == "__main__":
    main()
