# -*- coding: utf-8 -*-
"""Build the B2B Iguatemi break-even Excel model, styled like GE Beauty_BP_v2026.xlsx.

COGS split into two lines (Lucas 29/06):
  - Produto: % do sell-out (22,5%)
  - fee Boniteca: % do sell-in (8%)  -> aplicado sobre o sell-in, não sobre o sell-out.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

OUT = r"c:\claude\gebeauty\B2B_Breakeven_Iguatemi.xlsx"

# ---- palette / fonts (from BP: Calibri, brand red DF3630, red-font inputs) ----
RED   = "FFDF3630"; WHITE = "FFFFFFFF"; INPUTC= "FFFFF2CC"; GREY = "FFF2F2F2"
hdr_fill = PatternFill("solid", fgColor=RED)
in_fill  = PatternFill("solid", fgColor=INPUTC)
tot_fill = PatternFill("solid", fgColor=GREY)
f_hdr  = Font(name="Calibri", size=10, bold=True, color=WHITE)
f_lbl  = Font(name="Calibri", size=9)
f_lblb = Font(name="Calibri", size=9, bold=True)
f_in   = Font(name="Calibri", size=9, bold=True, color="FFC00000")
f_out  = Font(name="Calibri", size=9)
f_tot  = Font(name="Calibri", size=10, bold=True)
f_note = Font(name="Calibri", size=8, italic=True, color="FF808080")
thin   = Side(style="thin", color="FFBFBFBF")
box    = Border(left=thin, right=thin, top=thin, bottom=thin)

RS  = 'R$ #,##0;[Red]-R$ #,##0'
RS2 = 'R$ #,##0.00'
P0  = '0%'; P2 = '0.00%'; P1 = '0.0%'; INT = '#,##0'

wb = openpyxl.Workbook(); ws = wb.active
ws.title = "Break-even Iguatemi"
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 115
for c, w in {"A":2,"B":36,"C":15,"D":2,"E":2,"F":2,"G":3,"H":16,
             "I":12,"J":12,"K":12,"L":12,"M":12}.items():
    ws.column_dimensions[c].width = w

def put(cell, val, font=f_out, nf=None, fill=None, align=None, border=None):
    ws[cell] = val; ws[cell].font = font
    if nf: ws[cell].number_format = nf
    if fill: ws[cell].fill = fill
    if align: ws[cell].alignment = Alignment(horizontal=align, vertical="center")
    if border: ws[cell].border = border

def header(rng, text):
    ws.merge_cells(rng)
    put(rng.split(":")[0], text, font=f_hdr, fill=hdr_fill, align="left")
    for row in ws[rng]:
        for cell in row: cell.fill = hdr_fill

def inp(cell, val, nf):
    put(cell, val, font=f_in, nf=nf, fill=in_fill, align="right", border=box)

def out(cell, formula, nf=RS, font=f_out, fill=None, border=None):
    put(cell, formula, font=font, nf=nf, fill=fill, align="right", border=border)

# ---------------- title ----------------
header("B1:M1", "GE BEAUTY  |  Break-even B2B  —  Drogaria Iguatemi (Grupo DI)")
ws["B1"].font = Font(name="Calibri", size=13, bold=True, color=WHITE)
ws.row_dimensions[1].height = 22
ws.merge_cells("B2:M2")
put("B2", "Edite as células destacadas (fonte vermelha). Tudo recalcula. Premissas e ordem de DRE seguem o BP 2026.",
    font=f_note, align="left")

# ---------------- PREMISSAS (B:C, rows 5-20) ----------------
header("B4:C4", "PREMISSAS")
prem = [
 ("Nº de lojas (SP)",                       7,      INT),  # C5
 ("Horizonte (meses)",                      12,     INT),  # C6
 ("Sell-out / loja / mês",                  7000,   RS),   # C7
 ("Margem do varejista",                    0.40,   P0),   # C8
 ("Preço médio de tabela / un.",            96.34,  RS2),  # C9
 ("Nº de SKUs no portfólio",                15,     INT),  # C10
 ("Produto (% do sell-out)",                0.225,  P1),   # C11
 ("fee Boniteca (% do sell-in)",            0.08,   P1),   # C12
 ("ICMS efetivo (COMPETE-ES)",              0.011,  P2),   # C13
 ("PIS/COFINS (revenda monofásica)",        0.0,    P2),   # C14  reseller of monofásico = 0%
 ("Frete / pedido (ao CD)",                 500,    RS),   # C15
 ("Pedidos / ano",                          52,     INT),  # C16
 ("Investimento trade / ano",               150000, RS),   # C17
 ("Bonif.: unid / SKU / loja",              6,      INT),  # C18
 ("Bonif.: testers / SKU / loja",           1,      INT),  # C19
 ("Ativação / ano (evento+vitrine+seeding)",0,      RS),   # C20
]
for i, (label, val, nf) in enumerate(prem):
    r = 5 + i
    put(f"B{r}", label, font=f_lbl, align="left")
    inp(f"C{r}", val, nf)

# ---------------- DRE ANUAL (B:C, header 22, rows 23-35) ----------------
header("B22:C22", "DRE ANUAL DO CANAL (R$)")
dre = [
 ("Sell-out bruto (consumidor)",    "=C7*C5*C6",                  f_lbl),   # C23
 ("Receita bruta (sell-in)",        "=C7*(1-C8)*C5*C6",           f_lblb),  # C24
 ("(-) ICMS",                       "=-C13*C24",                  f_lbl),   # C25
 ("(-) PIS/COFINS",                 "=-C14*C24",                  f_lbl),   # C26
 ("Receita líquida",                "=C24+C25+C26",               f_lblb),  # C27
 ("(-) CMV — Produto",              "=-C11*C23",                  f_lbl),   # C28  Produto% x sell-out
 ("(-) fee Boniteca",               "=-C12*C24",                  f_lbl),   # C29  fee% x sell-in
 ("Margem bruta",                   "=C27+C28+C29",               f_lblb),  # C30
 ("(-) Frete logística",            "=-C15*C16",                  f_lbl),   # C31
 ("(-) Bonificação (1º pedido)",    "=-(C18+C19)*C10*C5*C9*C11",  f_lbl),   # C32  custo = Produto% (sem fee)
 ("(-) Investimento trade",         "=-C17",                      f_lbl),   # C33
 ("(-) Ativação",                   "=-C20",                      f_lbl),   # C34
 ("= RESULTADO DO CANAL",           "=C30+C31+C32+C33+C34",       f_tot),   # C35
]
for i, (label, formula, fnt) in enumerate(dre):
    r = 23 + i
    is_tot = "RESULTADO" in label
    put(f"B{r}", label, font=(f_tot if is_tot else fnt), align="left")
    out(f"C{r}", formula, nf=RS, font=(f_tot if is_tot else f_out),
        fill=(tot_fill if is_tot else None), border=(box if is_tot else None))
    if is_tot:
        ws[f"B{r}"].fill = tot_fill; ws[f"B{r}"].border = box

# ---------------- INDICADORES & BREAK-EVEN (B:C, header 37, rows 38-42) ----------------
header("B37:C37", "INDICADORES & BREAK-EVEN")
ind = [
 ("Margem bruta GE (% do sell-in)",     "=1-C11/(1-C8)-C12",                                          P1),  # C38
 ("Margem de contribuição (% sell-in)", "=1-C11/(1-C8)-C12-C13-C14",                                  P1),  # C39
 ("Sell-in / loja / mês",               "=C7*(1-C8)",                                                 RS),  # C40
 ("Break-even: sell-out / loja / mês",  "=(C15*C16+(C18+C19)*C10*C5*C9*C11+C17+C20)/(C5*C6*((1-C8)*(1-C13-C14-C12)-C11))", RS),  # C41
 ("Break-even: investim. trade máx.",   "=C30+C31+C32+C34",                                           RS),  # C42
]
for i, (label, formula, nf) in enumerate(ind):
    r = 38 + i
    put(f"B{r}", label, font=f_lblb, align="left", fill=tot_fill)
    out(f"C{r}", formula, nf=nf, font=f_lblb, fill=tot_fill, border=box)

# ---------------- SENSIBILIDADE grid (H:M) ----------------
# Eixo X = Investimento trade, ancorado em C17 (pedido deles, R$150k); margem FIXA = C8.
# Linha mediana H9 = sell-out de C7. COGS = Produto(C11, % sell-out) + Boniteca(C12, % sell-in).
header("H4:M4", "SENSIBILIDADE — Resultado anual (margem fixa = C8)")
put("H5", "sell-out ↓ / investimento →", font=f_lblb, align="left")
INV_STEP = 25000; SO_STEP = 1000

def axis(anchor, offset):
    return f"={anchor}" if offset == 0 else f"={anchor}{'+' if offset>0 else '-'}{abs(offset)}"

for j in range(5):                       # colunas I..M = investimento (mediana K5 = C17)
    col = get_column_letter(9 + j)
    put(f"{col}5", axis("$C$17", (j - 2) * INV_STEP),
        font=f_lblb, nf=RS, fill=tot_fill, align="right", border=box)
for i in range(7):                       # linhas H6..H12 = sell-out (mediana H9 = C7)
    r = 6 + i
    put(f"H{r}", axis("$C$7", (i - 3) * SO_STEP),
        font=f_lblb, nf=RS, fill=tot_fill, align="right", border=box)
    for j in range(5):
        col = get_column_letter(9 + j); invcol = f"{col}$5"
        # Resultado = sellout*L*H*((1-m)(1-icms-pis-boniteca)-produto) - frete - bonif - INVEST - ativ
        f = (f"=$H{r}*$C$5*$C$6*((1-$C$8)*(1-$C$13-$C$14-$C$12)-$C$11)"
             f"-$C$15*$C$16-($C$18+$C$19)*$C$10*$C$5*$C$9*$C$11-{invcol}-$C$20")
        out(f"{col}{r}", f, nf=RS, border=box)

grid_rng = "I6:M12"
ws.conditional_formatting.add(grid_rng, CellIsRule(operator="greaterThanOrEqual", formula=["0"],
    fill=PatternFill("solid", fgColor="FFC6EFCE"), font=Font(color="FF006100")))
ws.conditional_formatting.add(grid_rng, CellIsRule(operator="lessThan", formula=["0"],
    fill=PatternFill("solid", fgColor="FFFFC7CE"), font=Font(color="FF9C0006")))

# ---------------- notes ----------------
put("H15", "Notas", font=f_lblb, align="left")
notes = [
 "• Sell-in = sell-out × (1 − margem do varejista). GE fatura ao CD; varejista distribui.",
 "• COGS em 2 linhas: Produto = % do sell-out; fee Boniteca = % do SELL-IN.",
 "• Margem bruta GE (% sell-in) = 1 − Produto%/(1−margem) − Boniteca%.",
 "• Bonificação = 1º pedido grátis; custo = Produto% (sem fee Boniteca, pois não há venda).",
 "• Grade: margem fixa em C8; eixo X = investimento ao redor de C17 (passo R$25k);",
 "   linha do meio (H9) = seu sell-out de C7 (passo R$1k). Verde = lucro, vermelho = perda.",
 "• PIS/COFINS = 0%: GE REVENDE produtos monofásicos (cosméticos); o tributo é pago na",
 "   indústria (fornecedor) e já está embutido no custo do Produto. Sem crédito, sem débito.",
 "• ICMS 1,1% = COMPETE-ES atacadista (revenda). Flag p/ contabilidade: ICMS-ST.",
 "• Pedido deles (não acordado): 45% de margem + R$150.000 de trade.",
]
for i, t in enumerate(notes):
    put(f"H{16+i}", t, font=f_note, align="left")

ws.freeze_panes = "B3"
wb.save(OUT)
print("SAVED:", OUT)
