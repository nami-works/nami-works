"""Unilog B2C freight engine.

Parses the signed Unilog B2C workbook (freight matrix + CEP->GEOCOM-zone ranges)
and exposes freight(cep, weight_kg, nf_value) -> {base, gris, advalorem, total, zone, uf}.

Workbook lives in the shared Drive (root Unilog folder). Path overridable via
UNILOG_B2C_XLSX env var (used when the accented Drive path is unreadable and a
local ASCII copy is staged instead).

Freight per order (Anexo II.I):
  base       = matrix[zone][weight_bracket]  (R$; brackets are (lower, upper] kg,
               ceilings 0.25..30; above 30kg add R$/kg 'adicional' on the excess)
  gris       = GRIS%[zone]      * nf_value
  advalorem  = AdValorem%[zone] * nf_value
  total      = base + gris + advalorem
Tax (ISS+PIS/COFINS 14.25%) is treated as already embedded in the tariff (locked
premise #8). Returns uplift (+5%) is applied by the caller, not here.
"""
import os
from bisect import bisect_left
from pathlib import Path

import openpyxl

DEFAULT_XLSX = ("/g/Drives compartilhados/GEB_Financeiro/Contratos fornecedores/"
                "Unilog/GE-BEAUTY_B2C - Cosméticos - ORIGEM ES.xlsx")

# bracket ceilings from Tabela header row 1 (cols 1..34)
BRACKET_CEILINGS = [0.25, 0.30, 0.50, 0.75, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12,
                    13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30]


def _xlsx_path():
    return os.environ.get("UNILOG_B2C_XLSX", DEFAULT_XLSX)


def load():
    """Return (matrix, cep_index). matrix[zone] = dict; cep_index for lookup."""
    wb = openpyxl.load_workbook(_xlsx_path(), read_only=True, data_only=True)

    # ---- Tabela: freight matrix ----
    # Col layout (0-idx): 0=zone, 1..34=bracket rates (ceilings 0.25..30),
    # 35=adicional/kg (>30kg), 36=GRIS%, 37=AdValorem%.
    # Workbook typo: the 2nd 'SC/Capital 1' row is really 'SC/Interior 3'
    # (its 30kg rate ~R$592 is interior-3 magnitude; CEP ranges reference
    # 'SC/Interior 3', which has no own matrix row). Remap the duplicate.
    matrix = {}
    rows = list(wb["Tabela"].iter_rows(values_only=True))
    for r in rows[2:]:
        zone = r[0]
        if not zone:
            continue
        zone = zone.strip()
        if zone == "SC/Capital 1" and zone in matrix:
            zone = "SC/Interior 3"
        matrix[zone] = {
            "rates": list(r[1:35]),                    # 34 bracket rates
            "adicional_kg": r[35] or 0.0,
            "gris": r[36] or 0.0,
            "advalorem": r[37] or 0.0,
        }

    # ---- Abrangência: CEP ranges -> zone ----
    starts, ends, zones, ufs = [], [], [], []
    for r in wb["Abrangência"].iter_rows(min_row=2, values_only=True):
        ci, cf, cidade, uf, zone = r[0], r[1], r[2], r[3], r[4]
        if ci is None or cf is None or not zone:
            continue
        starts.append(int(ci)); ends.append(int(cf))
        zones.append(str(zone).strip()); ufs.append(str(uf).strip() if uf else None)

    order = sorted(range(len(starts)), key=lambda i: starts[i])
    cep_index = {
        "starts": [starts[i] for i in order],
        "ends":   [ends[i] for i in order],
        "zones":  [zones[i] for i in order],
        "ufs":    [ufs[i] for i in order],
    }
    return matrix, cep_index


def cep_to_int(cep):
    d = "".join(ch for ch in str(cep) if ch.isdigit())
    return int(d) if d else None


def lookup_zone(cep_index, cep, snap=True):
    """Return (zone, uf, exact). If the CEP falls in a gap between ranges and
    snap=True, snap to the nearest range (gaps are mostly xxxxxx0/1 boundary
    artifacts). exact=False flags a snapped result."""
    c = cep_to_int(cep)
    if c is None:
        return None, None, False
    starts, ends = cep_index["starts"], cep_index["ends"]
    i = bisect_left(starts, c + 1) - 1   # rightmost start <= c
    if 0 <= i and ends[i] >= c:
        return cep_index["zones"][i], cep_index["ufs"][i], True
    if not snap:
        return None, None, False
    # nearest by distance to [start,end] of candidates i and i+1
    cands = [j for j in (i, i + 1) if 0 <= j < len(starts)]
    if not cands:
        return None, None, False
    def dist(j):
        if c < starts[j]:
            return starts[j] - c
        if c > ends[j]:
            return c - ends[j]
        return 0
    j = min(cands, key=dist)
    return cep_index["zones"][j], cep_index["ufs"][j], False


def bracket_idx(weight_kg):
    for i, ceil in enumerate(BRACKET_CEILINGS):
        if weight_kg <= ceil:
            return i
    return len(BRACKET_CEILINGS) - 1   # >30kg handled via adicional


def freight(matrix, cep_index, cep, weight_kg, nf_value):
    zone, uf, exact = lookup_zone(cep_index, cep)
    if zone is None or zone not in matrix:
        return None
    m = matrix[zone]
    w = max(weight_kg or 0.0, 0.0)
    idx = bracket_idx(w)
    base = m["rates"][idx] or 0.0
    if w > 30 and m["adicional_kg"]:
        base = (m["rates"][-1] or 0.0) + (w - 30) * m["adicional_kg"]
    gris = m["gris"] * (nf_value or 0.0)
    adv = m["advalorem"] * (nf_value or 0.0)
    return {"zone": zone, "uf": uf, "exact": exact, "base": base, "gris": gris,
            "advalorem": adv, "total": base + gris + adv}


if __name__ == "__main__":
    matrix, idx = load()
    print(f"zones in matrix: {len(matrix)}")
    print(f"CEP ranges: {len(idx['starts'])}")
    zones_in_ranges = set(idx["zones"])
    zones_in_matrix = set(matrix)
    print(f"zones in ranges: {len(zones_in_ranges)}")
    print(f"in ranges but NOT in matrix: {sorted(zones_in_ranges - zones_in_matrix)}")
    print(f"in matrix but NOT in ranges: {sorted(zones_in_matrix - zones_in_ranges)}")
    # contiguity / overlap check
    gaps = overlaps = 0
    s, e = idx["starts"], idx["ends"]
    for i in range(1, len(s)):
        if s[i] <= e[i - 1]:
            overlaps += 1
        elif s[i] > e[i - 1] + 1:
            gaps += 1
    print(f"overlaps: {overlaps}  gaps: {gaps}  (total ranges {len(s)})")
    # spot checks
    for cep, w, v in [("01000-001", 0.5, 150), ("29173-795", 0.3, 100),
                      ("69900-000", 1.0, 200), ("70000-000", 0.75, 120)]:
        print(cep, w, v, "->", freight(matrix, idx, cep, w, v))
