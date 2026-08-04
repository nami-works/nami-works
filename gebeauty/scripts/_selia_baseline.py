"""Baseline current (Sélia, origin Extrema/MG) freight cost vs the Unilog table.

Parses the Sélia monthly 'Fretes' workbooks (per-shipment CT-e/Correios detail:
VALOR FRETE, VALOR NF, PESO, ESTADO) and computes current freight-cost intensity
overall and per macro-region, to decide apply-now vs wait-for-Unilog-cutover.

Read-only. Prints comparison; writes selia_baseline.json.
"""
import glob, json, os
from collections import defaultdict
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
DIR = HERE.parent/"logistics"/"frete-unilog"/"selia-actuals"
REGION={"N":["AC","AP","AM","PA","RO","RR","TO"],"NE":["AL","BA","CE","MA","PB","PE","PI","RN","SE"],
        "CO":["DF","GO","MT","MS"],"SE":["ES","MG","RJ","SP"],"S":["PR","RS","SC"]}
U2R={u:r for r,us in REGION.items() for u in us}
# Unilog gross freight % of revenue per region (from freight_final.json meta) for side-by-side
UNILOG_GROSS={"SE":0.122,"CO":0.139,"S":0.167,"NE":0.183,"N":0.286,"BR":0.139}


def num(v):
    try: return float(v)
    except (TypeError,ValueError): return None


def hdr_map(row):
    return {str(c).strip().upper(): i for i,c in enumerate(row) if c is not None}


def parse_sheet(ws):
    """Yield (uf, freight, nf, weight) from a carrier sheet by header names."""
    rows = ws.iter_rows(values_only=True)
    try: header = next(rows)
    except StopIteration: return
    h = hdr_map(header)
    ci_f = next((h[k] for k in h if "VALOR FRETE" in k), None)
    ci_nf = next((h[k] for k in h if k=="VALOR NF" or "VALOR NF" in k), None)
    ci_uf = next((h[k] for k in h if k in ("ESTADO","UF")), None)
    ci_w = next((h[k] for k in h if k=="PESO"), None)
    if ci_f is None: return
    for r in rows:
        if ci_f>=len(r): continue
        fr=num(r[ci_f])
        if fr is None or fr<=0: continue
        uf=(str(r[ci_uf]).strip().upper() if ci_uf is not None and ci_uf<len(r) and r[ci_uf] else None)
        nf=num(r[ci_nf]) if ci_nf is not None and ci_nf<len(r) else None
        w=num(r[ci_w]) if ci_w is not None and ci_w<len(r) else None
        yield uf, fr, nf, w


def main():
    files=sorted(glob.glob(str(DIR/"selia-fretes-2026-*.xlsx")))
    all_ship=[]
    per_file={}
    for f in files:
        wb=openpyxl.load_workbook(f, read_only=True, data_only=True)
        ship=[]
        for sn in wb.sheetnames:
            if sn in ("Descritivo","Desconto","GNRE - Postos Fiscais"): continue
            for rec in parse_sheet(wb[sn]):
                ship.append((sn,)+rec)
        per_file[os.path.basename(f)]={"n":len(ship),
            "freight":sum(x[2] for x in ship),
            "nf":sum(x[3] for x in ship if x[3])}
        all_ship+=ship

    tot_fr=sum(x[2] for x in all_ship)
    tot_nf=sum(x[3] for x in all_ship if x[3])
    with_uf=[x for x in all_ship if x[1] in U2R]
    no_uf=[x for x in all_ship if x[1] not in U2R]
    print("Files:")
    for k,v in per_file.items():
        print(f"  {k}: {v['n']} shipments, frete R${v['freight']:,.0f}, NF R${v['nf']:,.0f}")
    print(f"\nTOTAL: {len(all_ship)} shipments | frete R${tot_fr:,.0f} | NF R${tot_nf:,.0f}")
    print(f"shipments w/ UF: {len(with_uf)} ({100*len(with_uf)/len(all_ship):.0f}%)  "
          f"no-UF (Correios/intermpal): {len(no_uf)}")
    print(f"\nCURRENT freight / NF value (all): {100*tot_fr/tot_nf:.2f}%")
    print(f"avg freight/shipment: R${tot_fr/len(all_ship):.2f}")

    # per region (only UF-tagged shipments)
    reg=defaultdict(lambda:{"n":0,"fr":0.0,"nf":0.0})
    for sn,uf,fr,nf,w in with_uf:
        d=reg[U2R[uf]]; d["n"]+=1; d["fr"]+=fr; d["nf"]+=nf or 0
    print(f"\n{'Reg':<5}{'ship':>7}{'frete R$':>11}{'NF R$':>12}{'frt/NF':>8}{'avgFrt':>8}{'Unilog gross':>13}")
    reg_out={}
    for r in ["SE","CO","S","NE","N"]:
        d=reg[r]
        if not d["n"]: continue
        ratio=d["fr"]/d["nf"] if d["nf"] else 0
        print(f"{r:<5}{d['n']:>7}{d['fr']:>11,.0f}{d['nf']:>12,.0f}{100*ratio:>7.1f}%{d['fr']/d['n']:>8.2f}{100*UNILOG_GROSS[r]:>12.1f}%")
        reg_out[r]={"ship":d["n"],"freight":d["fr"],"nf":d["nf"],"frt_nf":ratio,"avg_freight":d["fr"]/d["n"],
                    "unilog_gross":UNILOG_GROSS[r]}
    uf_ratio=sum(d["fr"] for d in reg.values())/sum(d["nf"] for d in reg.values())
    print(f"\nUF-tagged freight/NF: {100*uf_ratio:.2f}%  (Unilog gross national {100*UNILOG_GROSS['BR']:.1f}%)")

    json.dump({"total_shipments":len(all_ship),"total_freight":tot_fr,"total_nf":tot_nf,
               "freight_nf_pct":tot_fr/tot_nf,"avg_freight":tot_fr/len(all_ship),
               "per_region":reg_out,"per_file":per_file,
               "no_uf_shipments":len(no_uf)}, open(HERE/"selia_baseline.json","w"), indent=2, default=float)
    print("\nwrote selia_baseline.json")


if __name__=="__main__":
    main()
