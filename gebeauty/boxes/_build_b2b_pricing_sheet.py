"""
Build the B2B pricing reference sheet:
  SKU | Product Title | Volume | Minimum Price (B2B break-even) | Retail Price | Flag

Sources:
  - Minimum prices from 'Estudo estoque 2026 apresentação (1).xlsx' (Planilha1).
  - Titles, retail prices, and volume from Shopify (live, pulled via MCP).
  - Volume = product metafield custom.dosagem.
  - Retail uses compareAtPrice when set, else price.
  - Filter: only productType="product" (excludes kit, acessorio, gift-card, rappi).

Output: 'B2B Pricing 2026.xlsx' (same folder).
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
INPUT_FILE = HERE / "Estudo estoque 2026  apresentação (1).xlsx"
OUTPUT_FILE = HERE / "B2B Pricing 2026.xlsx"

# Minimum (B2B break-even) prices keyed by current Shopify SKU.
# Source: Planilha1 of the input file. Dual-SKU rows mapped to the live SKU.
MIN_PRICE = {
    "GEB 001": 21.61,
    "GEB 008": 13.21,
    "GEB 002": 17.43,
    "GEB 003": 15.48,
    "GEB 023": 20.49,   # row 'GEB 004/023'
    "GEB 021": 17.29,   # row 'GEB 005/021'
    "GEB 020": 17.61,   # row 'GEB 006/020'
    "GEB 019": 17.80,   # row 'GEB 007/019'
    "GEB 022": 16.09,
    "GEB 013": 8.87,
    "GEB 010": 10.01,
    "GEB 011": 8.21,
    "GEB 101": 20.83,
    "GEB 102": 24.93,
    "GEB 120": 30.29,
    "GEB 024": 24.75,
    "GEB 121": 31.88,
    "GEB 029": 20.733,
}

# Active Shopify products with productType="product" only (pulled 2026-05-22).
# Excluded: productType "kit" (bundles incl. duplas), "acessorio", "gift-card",
# and "rappi" channel duplicates.
# Volume from product metafield custom.dosagem.
# (sku, title, volume, price, compareAtPrice or None)
SHOPIFY = [
    ("GEB 008", "shampoo a seco",                              "150mL",  69.00, None),
    ("GEB 019", "booster fortificante",                         "15mL",  75.00, None),
    ("GEB 003", "leave-in com proteção térmica",               "150mL",  99.00, None),
    ("GEB 002", "máscara condicionadora",                      "200mL",  95.00, None),
    ("GEB 001", "shampoo sem sulfato",                         "250mL",  95.00, None),
    ("GEB 020", "booster hidratante",                           "15mL",  69.00, None),
    ("GEB 021", "booster definição",                            "15mL",  69.00, None),
    ("GEB 023", "booster antioxidante",                         "15mL",  75.00, None),
    ("GEB 013", "travel size | shampoo sem sulfato",            "60mL",  40.00, None),
    ("GEB 011", "travel size | leave-in com proteção térmica",  "50mL",  47.00, None),
    ("GEB 010", "travel size | máscara condicionadora",         "50mL",  40.00, None),
    ("GEB 102", "primer liso intacto",                         "150mL", 139.00, None),
    ("GEB 101", "primer cachos definidos",                     "250mL", 149.00, None),
    ("GEB 022", "booster antifrizz",                            "15mL",  79.00, None),
    ("GEB 120", "leave-in pluma",                              "200mL", 149.00, None),
    ("GEB 024", "melon mood | body & hair splash",             "200mL", 129.00, None),
    ("GEB 121", "máscara mayday",                              "200g",  139.00, None),
    ("GEB 029", "travel size | melon mood body & hair splash", "100mL",  79.00, None),
]


def sku_sort_key(sku: str):
    # Sort so that "GEB nnn" runs together numerically, kits "GEBK###" follow,
    # then anything else alphabetically. Pure-numeric SKUs sort by number.
    s = sku.strip()
    if s.startswith("GEB ") and s[4:].isdigit():
        return (0, int(s[4:]), s)
    if s.startswith("GEBK") and s[4:].replace("MIX", "").isdigit():
        rest = s[4:]
        return (1, int(rest) if rest.isdigit() else 9999, s)
    if s.isdigit():
        return (2, int(s), s)
    return (3, 0, s)


def main():
    rows = []
    for sku, title, volume, price, compare_at in SHOPIFY:
        retail = compare_at if compare_at else price
        min_price = MIN_PRICE.get(sku)
        flag = "" if sku in MIN_PRICE else "NOT IN INPUT SHEET"
        rows.append((sku, title, volume, min_price, retail, flag))

    rows.sort(key=lambda r: sku_sort_key(r[0]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "B2B Pricing"

    headers = ["SKU", "Product Title", "Volume", "Minimum Price (BRL)", "Retail Price (BRL)", "Flag"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="FFFFFF")
    flag_fill = PatternFill("solid", fgColor="FEF3C7")
    border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )

    for col_idx, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = border

    for sku, title, volume, min_price, retail, flag in rows:
        ws.append([sku, title, volume, min_price, retail, flag])

    last_row = ws.max_row
    for r in range(2, last_row + 1):
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="left")
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="left")
        ws.cell(row=r, column=4).number_format = 'R$ #,##0.00'
        ws.cell(row=r, column=5).number_format = 'R$ #,##0.00'
        if ws.cell(row=r, column=6).value:
            for c in range(1, 7):
                ws.cell(row=r, column=c).fill = flag_fill
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = border

    widths = {1: 14, 2: 50, 3: 12, 4: 22, 5: 22, 6: 24}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A2"

    wb.save(OUTPUT_FILE)
    print(f"Wrote {OUTPUT_FILE}")
    print(f"Total rows: {len(rows)}")
    flagged = [r for r in rows if r[5]]
    print(f"Flagged (active in Shopify, missing from input): {len(flagged)}")


if __name__ == "__main__":
    main()
