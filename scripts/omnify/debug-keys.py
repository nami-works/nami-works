"""Debug: compare a few xlsx names to what Shopify returns."""
import json
import urllib.request
import openpyxl
from pathlib import Path

ENV_PATH = Path(r"C:\Users\Lucas Guimarães\Desktop\nami-works\sandbox\gebeauty\.env")
SRC = Path(r"C:\Users\Lucas Guimarães\Desktop\Pedidos_2026_05_19_142420.xlsx")

env = {l.split("=",1)[0].strip(): l.split("=",1)[1].strip()
       for l in ENV_PATH.read_text(encoding="utf-8").splitlines() if "=" in l and not l.startswith("#")}

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
ws = wb["Pedidos"]
hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
col = hdr.index("Pedido origem")
sample_names = []
for r in ws.iter_rows(min_row=2, max_row=10, values_only=True):
    raw = r[col]
    sample_names.append((raw, type(raw).__name__))
print("Sample xlsx 'Pedido origem' values + types:")
for raw, typ in sample_names:
    print(f"  {raw!r:>16}  type={typ}")

# Pull the first sample from Shopify and see what 'name' looks like.
domain = env["SHOPIFY_SHOP_DOMAIN"]
token = env["SHOPIFY_ADMIN_ACCESS_TOKEN"]
version = env.get("SHOPIFY_API_VERSION", "2025-01")
first = str(sample_names[0][0]).strip()
q = f"name:{first}"
body = json.dumps({
    "query": "query($q:String!){orders(query:$q,first:5){edges{node{id name}}}}",
    "variables": {"q": q}
}).encode()
req = urllib.request.Request(f"https://{domain}/admin/api/{version}/graphql.json",
    data=body, headers={"Content-Type":"application/json","X-Shopify-Access-Token":token})
with urllib.request.urlopen(req, timeout=30) as r:
    resp = json.loads(r.read().decode("utf-8"))
print(f"\nShopify query for name:{first}:")
print(json.dumps(resp, indent=2))
