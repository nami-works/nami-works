"""
Read Pedidos_2026_05_19_142420.xlsx, look up each row's Shopify order by
'Pedido origem' (column D), fetch the current fulfillment location, and
write a new xlsx with the location appended as a new column.

Highlights rows where the order has been flipped away from CD Extrema
(meaning Omie issued NF but the physical order is no longer at Extrema —
those NFs likely shouldn't have been emitted).
"""
import json
import sys
import urllib.request
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font

ENV_PATH = Path(r"C:\Users\Lucas Guimarães\Desktop\nami-works\sandbox\gebeauty\.env")
SRC = Path(r"C:\Users\Lucas Guimarães\Desktop\Pedidos_2026_05_19_142420.xlsx")
DST = Path(r"C:\Users\Lucas Guimarães\Desktop\Pedidos_2026_05_19_142420_enriched.xlsx")

EXTREMA_LOCATION_ID = "gid://shopify/Location/105538257216"


def load_env(path):
    env = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        env[k.strip()] = v.strip()
    return env


def shopify_graphql(domain, version, token, query, variables):
    url = f"https://{domain}/admin/api/{version}/graphql.json"
    body = json.dumps({"query": query, "variables": variables}).encode("utf-8")
    req = urllib.request.Request(url, data=body, headers={
        "Content-Type": "application/json",
        "X-Shopify-Access-Token": token,
    })
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode("utf-8"))


def chunk(seq, n):
    seq = list(seq)
    for i in range(0, len(seq), n):
        yield seq[i:i + n]


env = load_env(ENV_PATH)
domain = env["SHOPIFY_SHOP_DOMAIN"]
version = env.get("SHOPIFY_API_VERSION", "2025-01")
token = env["SHOPIFY_ADMIN_ACCESS_TOKEN"]

# Step 1: read source xlsx, extract order names.
wb = openpyxl.load_workbook(SRC)
ws = wb["Pedidos"]
print(f"Source rows (incl. header): {ws.max_row}")

# Find columns by header name.
hdr = [c.value for c in ws[1]]
col_pedido_origem = hdr.index("Pedido origem") + 1  # 1-based
col_nf = hdr.index("NF") + 1

# Pull all order names from column D.
order_names_in_xlsx = []
for r in range(2, ws.max_row + 1):
    cell = ws.cell(row=r, column=col_pedido_origem).value
    if cell is None:
        order_names_in_xlsx.append(None)
        continue
    order_names_in_xlsx.append(str(cell).strip())

unique_names = sorted({n for n in order_names_in_xlsx if n})
print(f"Unique order names to look up: {len(unique_names)}")

# Step 2: batch query Shopify.
gql = """
query GetOrders($q: String!, $first: Int!, $after: String) {
  orders(query: $q, first: $first, after: $after) {
    edges {
      cursor
      node {
        id name
        displayFulfillmentStatus
        fulfillments(first: 10) {
          location { id name }
          status
          createdAt
        }
        fulfillmentOrders(first: 10) {
          nodes {
            id status
            assignedLocation { location { id name } }
          }
        }
      }
    }
    pageInfo { hasNextPage endCursor }
  }
}
"""

resolved = {}  # name -> { location_id, location_name, fulfillment_status, status_per_fo }
for batch in chunk(unique_names, 25):
    q = " OR ".join(f"name:{n}" for n in batch)
    cursor = None
    while True:
        resp = shopify_graphql(domain, version, token, gql, {"q": q, "first": 50, "after": cursor})
        if "errors" in resp:
            print("Shopify error:", resp["errors"], file=sys.stderr); sys.exit(1)
        for e in resp["data"]["orders"]["edges"]:
            n = e["node"]
            name = n["name"]
            fos = n.get("fulfillmentOrders", {}).get("nodes") or []
            fulfillments = n.get("fulfillments") or []

            # Pick location: prefer OPEN FO (still unfulfilled, current physical location);
            # fall back to fulfillments[0].location (where the package was shipped from).
            location_id = None
            location_name = None
            source = None
            chosen_open = next((f for f in fos if f.get("status") == "OPEN"), None)
            if chosen_open:
                loc = (chosen_open.get("assignedLocation") or {}).get("location") or {}
                location_id = loc.get("id"); location_name = loc.get("name")
                source = f"open FO {chosen_open.get('id','?')[-12:]}"
            elif fulfillments:
                # Already shipped — use fulfillment's source location.
                loc = (fulfillments[0].get("location") or {})
                location_id = loc.get("id"); location_name = loc.get("name")
                source = f"fulfillment status={fulfillments[0].get('status')}"
            elif fos:
                # Closed FOs only — last resort.
                loc = (fos[0].get("assignedLocation") or {}).get("location") or {}
                location_id = loc.get("id"); location_name = loc.get("name")
                source = f"closed FO {fos[0].get('id','?')[-12:]} status={fos[0].get('status')}"

            resolved[name] = {
                "location_id": location_id,
                "location_name": location_name,
                "fulfillment_status": n.get("displayFulfillmentStatus"),
                "source": source,
                "fulfillments_count": len(fulfillments),
                "fos_count": len(fos),
            }
        info = resp["data"]["orders"]["pageInfo"]
        if not info.get("hasNextPage"):
            break
        cursor = info.get("endCursor")
    print(f"  batch -> resolved {len(resolved)}/{len(unique_names)}")

print(f"\nResolved: {len(resolved)}  Missing: {len(unique_names) - len(resolved)}")
missing = [n for n in unique_names if n not in resolved]
if missing:
    print(f"  missing names: {missing[:20]}")
# DEBUG: dump first 3 resolved entries to see actual shape.
print("\nDEBUG sample resolved entries:")
for k in list(resolved.keys())[:3]:
    print(f"  {k!r} -> {resolved[k]}")

# Step 3: write output xlsx with appended columns.
# Append: 'Local atual (Shopify)' + 'Status Shopify' + 'Flipped do CD Extrema?'
new_col_loc = ws.max_column + 1
new_col_status = ws.max_column + 2
new_col_flip = ws.max_column + 3
ws.cell(row=1, column=new_col_loc, value="Local atual (Shopify)").font = Font(bold=True)
ws.cell(row=1, column=new_col_status, value="Status fulfilment Shopify").font = Font(bold=True)
ws.cell(row=1, column=new_col_flip, value="Flipped do CD Extrema?").font = Font(bold=True)

yellow = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")  # flipped
red = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")     # missing on Shopify
green = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")   # still at Extrema (consistent)

flipped_count = 0
extrema_count = 0
missing_count = 0
flipped_rows = []  # (excel_row, order_name, nf, new_location)

for r in range(2, ws.max_row + 1):
    name = order_names_in_xlsx[r - 2]
    if not name:
        continue
    info = resolved.get(name)
    nf = ws.cell(row=r, column=col_nf).value
    if not info or not info["location_id"]:
        ws.cell(row=r, column=new_col_loc, value="NÃO ENCONTRADO NO SHOPIFY")
        ws.cell(row=r, column=new_col_status, value="-")
        ws.cell(row=r, column=new_col_flip, value="?")
        for c in range(new_col_loc, new_col_flip + 1):
            ws.cell(row=r, column=c).fill = red
        missing_count += 1
        continue
    ws.cell(row=r, column=new_col_loc, value=info["location_name"])
    ws.cell(row=r, column=new_col_status, value=info["fulfillment_status"])
    if info["location_id"] == EXTREMA_LOCATION_ID:
        ws.cell(row=r, column=new_col_flip, value="NÃO (em Extrema)")
        for c in range(new_col_loc, new_col_flip + 1):
            ws.cell(row=r, column=c).fill = green
        extrema_count += 1
    else:
        ws.cell(row=r, column=new_col_flip, value="SIM — NF emitida em local diferente")
        for c in range(new_col_loc, new_col_flip + 1):
            ws.cell(row=r, column=c).fill = yellow
        flipped_count += 1
        flipped_rows.append((r, name, nf, info["location_name"]))

# Auto-size new columns roughly.
for c in (new_col_loc, new_col_status, new_col_flip):
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 32

wb.save(DST)

print()
print(f"=== Summary ===")
print(f"  Total data rows:                       {ws.max_row - 1}")
print(f"  Still at CD Extrema (consistent):      {extrema_count}")
print(f"  FLIPPED away from Extrema (NF flag):   {flipped_count}")
print(f"  Not found on Shopify (?):              {missing_count}")
print()
if flipped_rows:
    print(f"Rows flagged (NF emitted but order is no longer at CD Extrema):")
    for r, name, nf, loc in flipped_rows:
        print(f"  row {r:>3}  pedido_origem=#{name}  NF={nf}  local_atual={loc}")
print()
print(f"Saved -> {DST}")
