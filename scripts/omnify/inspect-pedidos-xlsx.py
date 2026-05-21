"""
Quick inspect of the Omie pedidos spreadsheet to figure out column layout.
"""
import openpyxl
from pathlib import Path

XLSX = Path(r"C:\Users\Lucas Guimarães\Desktop\Pedidos_2026_05_19_142420.xlsx")
wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
print(f"Sheets: {wb.sheetnames}")
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n=== Sheet '{sn}' (dim: {ws.max_row} rows x {ws.max_column} cols) ===")
    rows = list(ws.iter_rows(values_only=True, max_row=8))
    for i, r in enumerate(rows):
        print(f"  row {i}: {r}")
    # Header inference: first row likely.
    if rows:
        hdr = rows[0]
        print(f"\n  Header columns: {[(i, c) for i, c in enumerate(hdr) if c is not None][:30]}")
