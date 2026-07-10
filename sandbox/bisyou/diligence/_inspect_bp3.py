# -*- coding: utf-8 -*-
import openpyxl
PATH = r"G:\Drives compartilhados\GEB_Financeiro\Orçamento e Resultados\2026\GE Beauty_BP_v2026.xlsx"
wv = openpyxl.load_workbook(PATH, data_only=True)["E-comm"]
wf = openpyxl.load_workbook(PATH, data_only=False)["E-comm"]

# Key DRE rows -> label
rows = {4:"Faturamento bruto",6:"Receita Formulas",44:"Receita bruta",45:"Impostos s/ rec",
        46:"ICMS",47:"PIS/COFINS",49:"Receita liquida",50:"CMV",51:"Produto(CMV)",55:"Margem bruta",
        58:"Armazenagem",59:"Frete s/ vendas",60:"Taxas comerciais",64:"Midia online",69:"Afiliadas",
        77:"Influencia",84:"Margem contribuicao",86:"SG&A",88:"Salarios",97:"Contabilidade",
        118:"Sistemas",121:"Rateio corp",131:"EBITDA",138:"IRPJ",144:"Lucro liquido",190:"# pedidos"}
# columns: BN=2026-01 (66), BO=2026-02 (67), CA=2026 annual (79), CB=% (80)
cols = {"BN(2026-01)":66, "BO(2026-02)":67, "CA(2026 ann)":79, "CB(% rec)":80}
out=[]
for r,lab in rows.items():
    out.append(f"\nr{r} {lab}:")
    for cname,c in cols.items():
        f=wf.cell(row=r,column=c).value
        v=wv.cell(row=r,column=c).value
        vv = round(v,3) if isinstance(v,float) else v
        out.append(f"   {cname}: {f!r}  -> {vv}")
open(r"C:\Users\Lucas Guimarães\Desktop\nami-works\sandbox\bisyou\diligence\_bp_2026_formulas.txt","w",encoding="utf-8").write("\n".join(out))
print("done")
