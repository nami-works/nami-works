# -*- coding: utf-8 -*-
import openpyxl, json
F=r"c:\claude\sandbox\bisyou\diligence\GE Beauty_BP_v2026__+Bisyou.xlsx"
wv=openpyxl.load_workbook(F,data_only=True)["Macro"]

# quarter column groups (1-based col indices)
Q={"2026Q1":[60,61,62],"2026Q2":[63,64,65],"2026Q3":[66,67,68],"2026Q4":[69,70,71],
   "2027Q1":[79,80,81],"2027Q2":[82,83,84],"2027Q3":[85,86,87],"2027Q4":[88,89,90]}
def qsum(r,cols):
    s=0.0
    for c in cols:
        v=wv.cell(row=r,column=c).value
        if isinstance(v,(int,float)): s+=v
    return s
def rowq(r): return {q:qsum(r,cols) for q,cols in Q.items()}

# 1) locate summary subtotal rows + SG&A block
labels_want={"Receita bruta","Receita líquida","CMV","Margem bruta","Despesas de venda",
 "Despesas de marketing","Margem de contribuição","SG&A","EBITDA","Lucro líquido",
 "Financiamento estoque"}
loc={}
for r in range(1,wv.max_row+1):
    d=wv.cell(row=r,column=4).value
    if isinstance(d,str) and d.strip() in labels_want and d.strip() not in loc:
        loc[d.strip()]=r
print("SUBTOTAL ROWS:",loc)

# 2) SG&A line items = Despesa rows between 'SG&A' subtotal and 'EBITDA'
sga_row=loc.get("SG&A"); ebitda_row=loc.get("EBITDA")
print("SG&A block: rows",sga_row,"->",ebitda_row)
sga_items=[]
for r in range((sga_row or 92)+1, ebitda_row or 143):
    b=wv.cell(row=r,column=2).value; d=wv.cell(row=r,column=4).value
    if b=="Despesa" and isinstance(d,str) and d.strip():
        # skip sub-subtotals that are intra-Macro sums? include all line items
        sga_items.append((r,d.strip()))
print("\nSG&A LINE ITEMS (%d):"%len(sga_items))
for r,d in sga_items: print(f"  r{r}: {d}")

# 3) build data dict: summary lines quarterly + 2026FY (col BV=74)
out={"summary":{},"sga":{},"meta":{}}
for lab,r in loc.items():
    out["summary"][lab]={"q":rowq(r),"fy2026":wv.cell(row=r,column=74).value}
for r,d in sga_items:
    out["sga"][d]={"row":r,"q":rowq(r),"fy2026":wv.cell(row=r,column=74).value}
json.dump(out, open("_ge_macro_data.json","w",encoding="utf-8"), ensure_ascii=False, indent=1)
print("\nwrote _ge_macro_data.json")
# quick print FY2026 summary
print("\nGE 2026 FY summary:")
for lab,r in loc.items():
    print(f"  {lab}: {wv.cell(row=r,column=74).value}")
PY
