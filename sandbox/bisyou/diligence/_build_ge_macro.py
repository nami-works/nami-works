# -*- coding: utf-8 -*-
import openpyxl, json
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

F=r"C:\Users\Lucas Guimarães\Desktop\nami-works\sandbox\bisyou\diligence\GE Beauty_BP_v2026__+Bisyou.xlsx"
D=json.load(open(r"C:\Users\Lucas Guimarães\Desktop\nami-works\sandbox\bisyou\diligence\_ge_macro_data.json",encoding="utf-8"))
wb=openpyxl.load_workbook(F)

NAVY="11243F"; YEL="FFF3CD"; HEAD=Font(bold=True,color="FFFFFF"); BOLD=Font(bold=True)
M='#,##0'; P='0.0%'; INPUT=PatternFill("solid",fgColor=YEL)
def headrow(ws,r,n):
    for c in range(1,n+1):
        cell=ws.cell(row=r,column=c)
        if cell.value is not None: cell.font=HEAD; cell.fill=PatternFill("solid",fgColor=NAVY)

# ---------- TAB: GE_Macro + Savings ----------
if "GE_Macro_Savings" in wb.sheetnames: del wb["GE_Macro_Savings"]
ws=wb.create_sheet("GE_Macro_Savings")
ws["A1"]="GE BEAUTY — macro P&L (quarterly) + savings from sharing Ops/Fin overhead with Bisyou"; ws["A1"].font=BOLD

qcols=["2026Q1","2026Q2","2026Q3","2026Q4","2026 FY","2027Q1","2027Q2","2027Q3","2027Q4","2027 FY"]
ws["A3"]="GE P&L (as-is), R$"
for j,q in enumerate(qcols): ws.cell(row=3,column=2+j,value=q)
headrow(ws,3,11)
def fy(s,yr): return sum(D["summary"][s]["q"][f"{yr}Q{i}"] for i in (1,2,3,4))
plines=["Receita bruta","Receita líquida","CMV","Margem bruta","Despesas de venda",
        "Despesas de marketing","Margem de contribuição","SG&A","Financiamento estoque","EBITDA","Lucro líquido"]
prow={}
r=4
for lab in plines:
    disp = "Boniteca (Financ. estoque)" if lab=="Financiamento estoque" else lab
    ws.cell(row=r,column=1,value=disp)
    q=D["summary"][lab]["q"]
    vals=[q["2026Q1"],q["2026Q2"],q["2026Q3"],q["2026Q4"],fy(lab,"2026"),
          q["2027Q1"],q["2027Q2"],q["2027Q3"],q["2027Q4"],fy(lab,"2027")]
    for j,v in enumerate(vals): ws.cell(row=r,column=2+j,value=v).number_format=M
    if lab in ("Margem bruta","Margem de contribuição","EBITDA","Lucro líquido","Receita líquida"):
        for c in range(1,12): ws.cell(row=r,column=c).font=BOLD
    prow[lab]=r; r+=1
EBITDA_ROW=prow["EBITDA"]

# ----- savings section -----
r+=2
ws.cell(row=r,column=1,value="SAVINGS FROM SHARING OPS/FIN OVERHEAD WITH BISYOU").font=BOLD
r+=1
ws.cell(row=r,column=1,value="Bisyou's share of each shared overhead line  →")
share_cell=f"B{r}"
c=ws.cell(row=r,column=2,value=0.50); c.number_format=P; c.fill=INPUT; c.font=BOLD
ws.cell(row=r,column=3,value="(editable. 50% = split evenly between the two brands)")
r+=1
ws.cell(row=r,column=1,value="Savings begin H2-2026 (Q3), matching Bisyou onboarding; full-year in 2027.")
r+=2
hdr=r
ws.cell(row=r,column=1,value="SG&A category"); ws.cell(row=r,column=2,value="Share? (1/0)")
ws.cell(row=r,column=3,value="GE cost 2026 H2"); ws.cell(row=r,column=4,value="GE cost 2027 FY")
ws.cell(row=r,column=5,value="Savings 2026"); ws.cell(row=r,column=6,value="Savings 2027")
headrow(ws,r,6)
SUBTOTALS={"Time CLT","Sistemas","Escritório","Serviços contratados","Outras despesas"}
SHARED_DEFAULT={"Prestadores de Serviço","ERP","E-mail + ferramentas","Aluguel","Condomínio","IPTU",
 "Energia elétrica","Manutenção predial","Materiais de uso e consumo","Contabilidade","Advocacia",
 "Manutenção equipamentos","Assinaturas e filiações","Internet","Consultoria","Assessoria regulatória",
 "Sustentação e performance","Tarifas bancárias","Cartório","Reembolsos","Taxas legais","Locação equipamentos"}
r+=1
first_item=r
for cat,info in D["sga"].items():
    if cat in SUBTOTALS: continue
    q=info["q"]
    h2_2026=q["2026Q3"]+q["2026Q4"]
    fy2027=sum(q[f"2027Q{i}"] for i in (1,2,3,4))
    if abs(h2_2026)<1 and abs(fy2027)<1: continue  # skip empty lines
    ws.cell(row=r,column=1,value=cat)
    tg=ws.cell(row=r,column=2,value=1 if cat in SHARED_DEFAULT else 0); tg.fill=INPUT; tg.font=BOLD
    ws.cell(row=r,column=3,value=h2_2026).number_format=M
    ws.cell(row=r,column=4,value=fy2027).number_format=M
    ws.cell(row=r,column=5,value=f"=-C{r}*B{r}*${share_cell}").number_format=M
    ws.cell(row=r,column=6,value=f"=-D{r}*B{r}*${share_cell}").number_format=M
    r+=1
last_item=r-1
# totals
ws.cell(row=r,column=1,value="TOTAL SAVINGS FOR GE BEAUTY").font=BOLD
ws.cell(row=r,column=5,value=f"=SUM(E{first_item}:E{last_item})").number_format=M
ws.cell(row=r,column=6,value=f"=SUM(F{first_item}:F{last_item})").number_format=M
ws.cell(row=r,column=5).font=BOLD; ws.cell(row=r,column=6).font=BOLD
tot_row=r
r+=2
# adjusted EBITDA
ws.cell(row=r,column=1,value="GE EBITDA (as-is)").font=BOLD
ws.cell(row=r,column=5,value=fy("EBITDA","2026")).number_format=M
ws.cell(row=r,column=6,value=fy("EBITDA","2027")).number_format=M
r+=1
ws.cell(row=r,column=1,value="GE EBITDA + Ops/Fin savings (adjusted)").font=BOLD
ws.cell(row=r,column=5,value=f"=E{tot_row}+E{r-1}").number_format=M
ws.cell(row=r,column=6,value=f"=F{tot_row}+F{r-1}").number_format=M
for cc in (1,5,6): ws.cell(row=r,column=cc).font=BOLD
r+=2
ws.cell(row=r,column=1,value="Note: toggles default to a suggested Ops/Fin/back-office set — REVIEW each. Savings = GE cost × Share? × share%. Group-neutral (Bisyou absorbs the share); this view shows GE's standalone benefit.")
ws.column_dimensions['A'].width=34
for cl in "BCDEFGHIJK": ws.column_dimensions[cl].width=13

# ---------- Bisyou tab: benchmark columns ----------
b=wb["Bisyou"]
# GE 2026 FY %s of net revenue
nrev=D["summary"]["Receita líquida"]["fy2026"]
gepct={
 "Receita líquida":1.0,
 "CMV":D["summary"]["CMV"]["fy2026"]/nrev,
 "Margem bruta":D["summary"]["Margem bruta"]["fy2026"]/nrev,
 "Despesas de venda":D["summary"]["Despesas de venda"]["fy2026"]/nrev,
 "Marketing":D["summary"]["Despesas de marketing"]["fy2026"]/nrev,
 "SG&A":D["summary"]["SG&A"]["fy2026"]/nrev,
 "Boniteca":D["summary"]["Financiamento estoque"]["fy2026"]/nrev,
 "EBITDA":D["summary"]["EBITDA"]["fy2026"]/nrev,
}
# map Bisyou DRE rows (17-27) -> benchmark key
rowmap={17:"Receita líquida",18:"CMV",19:"Margem bruta",20:"Despesas de venda",21:"Marketing",
        22:"SG&A",24:"Boniteca",25:"EBITDA"}
b["E16"]="Bisyou % net"; b["F16"]="GE 2026 % (benchmark)"; b["E16"].font=BOLD; b["F16"].font=BOLD
for rr in range(17,28):
    # Bisyou % of net (conservative col C / net C4)
    if b.cell(row=rr,column=3).value is not None:
        b.cell(row=rr,column=5,value=f"=C{rr}/$C$4").number_format=P
    if rr in rowmap:
        b.cell(row=rr,column=6,value=gepct[rowmap[rr]]).number_format=P
b["E29"]="Benchmark: GE 2026 FY structure (% of net). Royalty has no GE equivalent. NB: GE EBITDA % is PRE-Boniteca; Bisyou EBITDA is post-Boniteca+royalty — not strictly like-for-like."
b.column_dimensions['E'].width=14; b.column_dimensions['F'].width=20

wb.save(F)
chk=openpyxl.load_workbook(F)
print("saved. sheets:",len(chk.sheetnames),"| charts:",sum(len(getattr(chk[s],"_charts",[])) for s in chk.sheetnames))
print("new tab present:", "GE_Macro_Savings" in chk.sheetnames)
print("Bisyou tweaks intact: C7(mkt)=",chk["Bisyou"]["C7"].value," C8(SGA)=",chk["Bisyou"]["C8"].value)
print("GE FY2026 %s:", {k:round(v,3) for k,v in gepct.items()})
