# -*- coding: utf-8 -*-
import openpyxl, unicodedata
from openpyxl.styles import Font, PatternFill
ORIG=r"G:\Drives compartilhados\GEB_Financeiro\Orçamento e Resultados\2026\GE Beauty_BP_v2026.xlsx"
COPY=r"c:\claude\sandbox\bisyou\diligence\GE Beauty_BP_v2026__+Bisyou.xlsx"
def norm(s):
    return unicodedata.normalize("NFKD",str(s)).encode("ascii","ignore").decode().strip().lower() if s is not None else ""
macv=openpyxl.load_workbook(ORIG,data_only=True)["Macro"]
wbf=openpyxl.load_workbook(COPY,data_only=False)
g=wbf["GE_Macro_Savings"]; b=wbf["Bisyou"]
mac_map={}
for r in range(1,macv.max_row+1):
    k=norm(macv.cell(row=r,column=4).value)
    if k: mac_map.setdefault(k,r)
Q26={1:[60,61,62],2:[63,64,65],3:[66,67,68],4:[69,70,71]}
Q27={1:[79,80,81],2:[82,83,84],3:[85,86,87],4:[88,89,90]}
def qs(r,cols):
    s=0.0;a=False
    for c in cols:
        v=macv.cell(row=r,column=c).value
        if isinstance(v,(int,float)): s+=v;a=True
    return s if a else None
def fetch(mr):
    q26=[qs(mr,Q26[i]) for i in(1,2,3,4)]; q27=[qs(mr,Q27[i]) for i in(1,2,3,4)]
    fy26=sum(v for v in q26 if v) if any(q26) else None
    fy27=sum(v for v in q27 if v) if any(q27) else None
    return q26+[fy26]+q27+[fy27]

YEL=PatternFill("solid",fgColor="FFF3CD"); BOLD=Font(bold=True); M='#,##0'

# 1) REFRESH top P&L rows 4-14 from pristine original (fix drift)
top={4:"receita bruta",5:"receita liquida",6:"cmv",7:"margem bruta",8:"despesas de venda",
 9:"despesas de marketing",10:"margem de contribuicao",11:"sg&a",12:"financiamento estoque",
 13:"ebitda",14:"lucro liquido"}
for rr,key in top.items():
    mr=mac_map.get(key)
    if mr:
        for j,v in enumerate(fetch(mr)):
            if v is not None: g.cell(row=rr,column=2+j,value=v).number_format=M

# 2) PER-ROW SAVINGS on the SG&A breakdown rows
g.cell(row=17,column=12,value="Share % -> Bisyou").font=Font(bold=True,color="FFFFFF"); g.cell(row=17,column=12).fill=PatternFill("solid",fgColor="11243F")
g.cell(row=17,column=13,value="Savings 2026 (H2)").font=Font(bold=True,color="FFFFFF"); g.cell(row=17,column=13).fill=PatternFill("solid",fgColor="11243F")
g.cell(row=17,column=14,value="Savings 2027 (FY)").font=Font(bold=True,color="FFFFFF"); g.cell(row=17,column=14).fill=PatternFill("solid",fgColor="11243F")
# default share% by row label (norm)
defshare={"financeiro":0.5,"operacoes":0.5,"comercial":0.0,"erp":0.5,"e-mail + ferramentas":0.5,
          "escritorio":0.5,"servicos contratados":0.5}
sav_rows=[]
for rr in range(18,50):
    lab=g.cell(row=rr,column=1).value
    k=norm(lab)
    if k in defshare:
        l=g.cell(row=rr,column=12,value=defshare[k]); l.number_format='0%'; l.fill=YEL; l.font=BOLD
        g.cell(row=rr,column=13,value=f"=-(D{rr}+E{rr})*L{rr}").number_format=M     # 2026 H2 = Q3+Q4
        g.cell(row=rr,column=14,value=f"=-K{rr}*L{rr}").number_format=M             # 2027 FY
        sav_rows.append(rr)
# total + adjusted EBITDA
tr=51
g.cell(row=tr,column=1,value="TOTAL OPS/FIN SAVINGS FOR GE BEAUTY").font=BOLD
g.cell(row=tr,column=13,value="=SUM(M18:M49)").number_format=M
g.cell(row=tr,column=14,value="=SUM(N18:N49)").number_format=M
g.cell(row=tr,column=13).font=BOLD; g.cell(row=tr,column=14).font=BOLD
g.cell(row=53,column=1,value="GE EBITDA (as-is)  →  FY2026 / FY2027").font=BOLD
g.cell(row=53,column=6,value="=F46").number_format=M
g.cell(row=53,column=11,value="=K46").number_format=M
g.cell(row=54,column=1,value="GE EBITDA + Ops/Fin savings (adjusted)").font=BOLD
g.cell(row=54,column=6,value="=F46+M51").number_format=M
g.cell(row=54,column=11,value="=K46+N51").number_format=M
g.cell(row=54,column=6).font=BOLD; g.cell(row=54,column=11).font=BOLD
g.cell(row=56,column=1,value="Per-row savings = -(GE cost) x Share%. Set Share% per line (yellow). Fill Financeiro/Operacoes (your split of Time) and their savings auto-compute. 2026 = H2 only (Bisyou onboards mid-year).")
g.column_dimensions['L'].width=14; g.column_dimensions['M'].width=15; g.column_dimensions['N'].width=15

# 3) REFRESH Bisyou benchmark GE %s from pristine original
nrev=fetch(mac_map["receita liquida"])[4]
def gepct(key):
    mr=mac_map.get(key);
    return (fetch(mr)[4]/nrev) if mr else None
bench={17:"receita liquida",18:"cmv",19:"margem bruta",20:"despesas de venda",21:"despesas de marketing",
       22:"sg&a",24:"financiamento estoque",25:"ebitda"}
for rr,key in bench.items():
    val=1.0 if key=="receita liquida" else gepct(key)
    if val is not None: b.cell(row=rr,column=6,value=val).number_format='0.0%'

wbf.save(COPY)
print("done. savings rows:",sav_rows)
print("Top P&L refreshed from pristine original.")
# verify reconciliation
chk=openpyxl.load_workbook(COPY,data_only=False)["GE_Macro_Savings"]
print("Top SG&A FY26 r11:",chk.cell(row=11,column=6).value," granular SG&A FY26 r36:",chk.cell(row=36,column=6).value)
print("Top EBITDA FY26 r13:",chk.cell(row=13,column=6).value," granular EBITDA FY26 r46:",chk.cell(row=46,column=6).value)
