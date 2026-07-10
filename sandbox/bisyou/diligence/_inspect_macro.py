# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.utils import get_column_letter
PATH = r"G:\Drives compartilhados\GEB_Financeiro\Orçamento e Resultados\2026\GE Beauty_BP_v2026.xlsx"
OUT  = r"c:\claude\sandbox\bisyou\diligence\_bp_macro_dump.txt"
wv = openpyxl.load_workbook(PATH, data_only=True)["Macro"]
wf = openpyxl.load_workbook(PATH, data_only=False)["Macro"]
L=[f"DIMS {wv.dimensions} max_row={wv.max_row} max_col={wv.max_column}={get_column_letter(wv.max_column)}"]
maxc=min(wv.max_column,40); maxr=min(wv.max_row,160)
for r in range(1,maxr+1):
    cells=[]
    for c in range(1,maxc+1):
        v=wv.cell(row=r,column=c).value; f=wf.cell(row=r,column=c).value
        if v is None and f is None: continue
        col=get_column_letter(c)
        if isinstance(f,str) and f.startswith("="):
            vv="" if v is None else (round(v,2) if isinstance(v,float) else (f"{v.year}-{v.month:02d}" if hasattr(v,"year") else v))
            cells.append(f"{col}{r}[{f} ->{vv}]")
        else:
            vv=round(v,3) if isinstance(v,float) else (f"{v.year}-{v.month:02d}" if hasattr(v,"year") else v)
            cells.append(f"{col}{r}={vv}")
    if cells: L.append("  ".join(cells))
open(OUT,"w",encoding="utf-8").write("\n".join(L))
print("wrote",OUT,"lines",len(L),"| sheet dims",wv.dimensions)
