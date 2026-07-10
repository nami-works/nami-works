# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"G:\Drives compartilhados\GEB_Financeiro\Orçamento e Resultados\2026\GE Beauty_BP_v2026.xlsx"
OUT = r"c:\claude\sandbox\bisyou\diligence\Bisyou_Impact_Model.xlsx"

# ---- 1) Read GE group 2026 totals (Macro col BV = 2026 annual) ----
mv = openpyxl.load_workbook(SRC, data_only=True)["Macro"]
BV = 74  # column BV
ge = {
 "Receita bruta": mv.cell(row=46,column=BV).value,
 "Receita líquida": mv.cell(row=51,column=BV).value,
 "Margem bruta": mv.cell(row=57,column=BV).value,
 "Margem de contribuição": mv.cell(row=91,column=BV).value,
 "EBITDA": mv.cell(row=143,column=BV).value,
 "Lucro líquido": mv.cell(row=161,column=BV).value,
}
print("GE group 2026 (Macro BV):")
for k,v in ge.items(): print(f"  {k}: {v}")

# ---- 2) Boniteca regressive rate (GE table, monthly gross formula-revenue) ----
THR=[500000,750000,1000000,1250000,1500000,2000000,10000000]
RATE=[0.10,0.09,0.08,0.075,0.065,0.06,0.05]
def boniteca_rate(monthly_gross):
    cg=sum(1 for t in THR if t>monthly_gross)
    idx=len(THR)-cg  # 0-based: COUNTA-COUNTIF+1 -> 1-based (8-cg); 0-based (7-cg)
    idx=max(0,min(idx,len(RATE)-1))
    return RATE[idx]

# ---- 3) Bisyou aggregate DRE ----
GM=0.80; SELL=0.17; MKT=0.25; SGA=1_250_000; TAXG=0.13  # gross-up: net = gross*(1-TAXG)
def bisyou(net_annual, royalty):
    gross=net_annual/(1-TAXG)
    bon_rate=boniteca_rate(gross/12)
    boniteca=bon_rate*gross
    gp=GM*net_annual
    sell=SELL*net_annual; mkt=MKT*net_annual; roy=royalty*net_annual
    ebitda=gp-sell-mkt-SGA-roy-boniteca
    return dict(net=net_annual,gross=gross,bon_rate=bon_rate,gp=gp,sell=sell,mkt=mkt,
                sga=SGA,roy=roy,boniteca=boniteca,ebitda=ebitda)

CONS=bisyou(5_500_000,0.08); RUN=bisyou(7_600_000,0.08)
print("\nBoniteca rate cons/run:",CONS['bon_rate'],RUN['bon_rate'])
print("EBITDA cons (5.5M,8pct): {:.0f} | run (7.6M,8pct): {:.0f}".format(CONS["ebitda"],RUN["ebitda"]))

# breakeven net @8% and royalty ceilings
def be_net(r): return SGA/(GM-SELL-MKT-r-boniteca_rate((6e6/(1-TAXG))/12)/(1-TAXG))  # approx using ~9% bon as %of net
def roy_ceiling(net):
    gross=net/(1-TAXG); bon=boniteca_rate(gross/12)*gross
    # ebitda=0 -> roy_amt = gp-sell-mkt-sga-bon ; r=roy_amt/net
    return (GM*net-SELL*net-MKT*net-SGA-bon)/net
print("Royalty ceiling cons:{:.3f} run:{:.3f}".format(roy_ceiling(5_500_000),roy_ceiling(7_600_000)))

# ---- 4) Build workbook ----
wb=openpyxl.Workbook()
NAVY="11243F"; GOLD="C79A3A"; HEAD=Font(bold=True,color="FFFFFF"); BOLD=Font(bold=True)
def hd(ws,cells_row):
    for c in ws[cells_row]:
        if c.value is not None: c.font=HEAD; c.fill=PatternFill("solid",fgColor=NAVY)
def money(ws,coord,val):
    ws[coord]=val; ws[coord].number_format='#,##0'
def pct(ws,coord,val):
    ws[coord]=val; ws[coord].number_format='0.0%'

# --- Tab 1: Bisyou channel DRE (both bases) ---
ws=wb.active; ws.title="Bisyou_DRE"
ws["A1"]="BISYOU — aggregate channel DRE (R$/yr run-rate). Boniteca via GE regressive table; royalty + Boniteca both ABOVE EBITDA."
ws["A1"].font=BOLD
ws["A3"]="Line"; ws["B3"]="% of net"; ws["C3"]="Conservative (R$5.5M net)"; ws["D3"]="Run-rate (R$7.6M net)"
hd(ws,3)
rows=[
 ("Receita bruta (gross)", None, CONS['gross'], RUN['gross']),
 ("Receita líquida (net)", 1.00, CONS['net'], RUN['net']),
 ("(–) CMV (20%)", -0.20, -(CONS['net']-CONS['gp']), -(RUN['net']-RUN['gp'])),
 ("= Margem bruta (80%)", 0.80, CONS['gp'], RUN['gp']),
 ("(–) Despesas de venda (17%)", -0.17, -CONS['sell'], -RUN['sell']),
 ("(–) Marketing / CPA (25%)", -0.25, -CONS['mkt'], -RUN['mkt']),
 ("(–) SG&A (lean, fixed)", None, -SGA, -SGA),
 ("(–) Royalty (8% of net)", -0.08, -CONS['roy'], -RUN['roy']),
 ("(–) Boniteca fee (GE table, ~%.0f%% of gross)"%(CONS['bon_rate']*100), None, -CONS['boniteca'], -RUN['boniteca']),
 ("= EBITDA", None, CONS['ebitda'], RUN['ebitda']),
]
r=4
for lab,p,cv,rv in rows:
    ws.cell(row=r,column=1,value=lab)
    if p is not None: ws.cell(row=r,column=2,value=p).number_format='0.0%'
    ws.cell(row=r,column=4-1,value=cv).number_format='#,##0'  # C
    ws.cell(row=r,column=4,value=rv).number_format='#,##0'    # D
    if lab.startswith("="):
        for c in (1,3,4): ws.cell(row=r,column=c).font=BOLD
    r+=1
ws["A16"]="Note: GE's Boniteca table is applied to Bisyou's smaller monthly volume → ~9% (vs Bisyou's own contract ~7%). Start = H2-2026 (Jul–Dec) = half of annual."
ws.column_dimensions['A'].width=42
for col in ('B','C','D'): ws.column_dimensions[col].width=20

# --- Tab 2: Group impact bridge ---
g=wb.create_sheet("Group_Impact")
g["A1"]="GROUP P&L IMPACT — GE 2026 (from BP Macro) + Bisyou H2-2026 (conservative). R$"; g["A1"].font=BOLD
g["A3"]="Line"; g["B3"]="GE 2026 (BP)"; g["C3"]="Bisyou H2-2026"; g["D3"]="Pro-forma group"; g["E3"]="Δ %"
hd(g,3)
half=0.5
bridge=[
 ("Receita líquida", ge["Receita líquida"], CONS['net']*half),
 ("Margem bruta", ge["Margem bruta"], CONS['gp']*half),
 ("Margem de contribuição", ge["Margem de contribuição"], (CONS['gp']-CONS['sell']-CONS['mkt'])*half),
 ("EBITDA", ge["EBITDA"], CONS['ebitda']*half),
]
r=4
for lab,gev,bv in bridge:
    g.cell(row=r,column=1,value=lab)
    g.cell(row=r,column=2,value=gev).number_format='#,##0'
    g.cell(row=r,column=3,value=bv).number_format='#,##0'
    g.cell(row=r,column=4,value=(gev or 0)+bv).number_format='#,##0'
    g.cell(row=r,column=5,value=(bv/gev) if gev else None)
    if g.cell(row=r,column=5).value is not None: g.cell(row=r,column=5).number_format='0.0%'
    r+=1
g["A10"]="Bisyou H2-2026 = half of conservative annual run-rate. EBITDA shown is Bisyou's (post-Boniteca, post-royalty). NB: GE's own EBITDA is PRE-Boniteca (model convention) — not strictly like-for-like."
g.column_dimensions['A'].width=26
for col in ('B','C','D','E'): g.column_dimensions[col].width=18

# --- Tab 3: Sensitivity grid + royalty ceiling ---
s=wb.create_sheet("Sensitivity")
s["A1"]="EBITDA CONTRIBUTION (R$/yr, full-year run-rate) — Royalty % (rows) × Net revenue (cols). Boniteca via GE table; SG&A R$1.25M; mktg 25%."; s["A1"].font=BOLD
revs=[5_000_000,5_500_000,6_000_000,6_500_000,7_000_000,7_600_000,8_000_000]
roys=[0.03,0.04,0.05,0.06,0.07,0.08]
s["A3"]="Royalty ↓ / Net rev →"
for j,rv in enumerate(revs):
    c=s.cell(row=3,column=2+j,value=rv); c.number_format='#,##0'; c.font=HEAD; c.fill=PatternFill("solid",fgColor=NAVY)
s["A3"].font=HEAD; s["A3"].fill=PatternFill("solid",fgColor=NAVY)
for i,ro in enumerate(roys):
    s.cell(row=4+i,column=1,value=ro).number_format='0%'
    s.cell(row=4+i,column=1).font=BOLD
    for j,rv in enumerate(revs):
        e=bisyou(rv,ro)['ebitda']
        cell=s.cell(row=4+i,column=2+j,value=e); cell.number_format='#,##0'
        if e<0: cell.font=Font(color="B23B3B")
        else: cell.font=Font(color="1F7A3F")
# royalty ceiling row
s.cell(row=12,column=1,value="Royalty ceiling (EBITDA=0)").font=BOLD
for j,rv in enumerate(revs):
    s.cell(row=12,column=2+j,value=roy_ceiling(rv)).number_format='0.0%'
s["A14"]="Read: at 8% royalty, Bisyou is accretive only above ~R$6.4M net. Royalty the deal bears: ~5% at R$5.5M, ~11% at R$7.6M. Green=accretive, red=dilutive."
s.column_dimensions['A'].width=24
for j in range(len(revs)): s.column_dimensions[get_column_letter(2+j)].width=14

wb.save(OUT)
print("\nSaved", OUT)
print("Breakeven net @8pct royalty ~ R$ {:.0f}".format((SGA)/(GM-SELL-MKT-0.08-CONS["boniteca"]/CONS["net"])))
PY
