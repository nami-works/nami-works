# -*- coding: utf-8 -*-
import openpyxl, copy
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

SRC=r"G:\Drives compartilhados\GEB_Financeiro\Orçamento e Resultados\2026\GE Beauty_BP_v2026.xlsx"
STD=r"C:\Users\Lucas Guimarães\Desktop\nami-works\sandbox\bisyou\diligence\Bisyou_Impact_Model.xlsx"
BPC=r"C:\Users\Lucas Guimarães\Desktop\nami-works\sandbox\bisyou\diligence\GE Beauty_BP_v2026__+Bisyou.xlsx"

mv=openpyxl.load_workbook(SRC,data_only=True)["Macro"]; BV=74
GE={"Receita líquida":mv.cell(row=51,column=BV).value,"Margem bruta":mv.cell(row=57,column=BV).value,
    "Margem de contribuição":mv.cell(row=91,column=BV).value,"EBITDA":mv.cell(row=143,column=BV).value}

NAVY="11243F"; YEL="FFF3CD"; HEAD=Font(bold=True,color="FFFFFF"); BOLD=Font(bold=True)
INPUT_FILL=PatternFill("solid",fgColor=YEL)
M='#,##0'; P='0.0%'

def style_head(ws,row,ncol):
    for c in range(1,ncol+1):
        cell=ws.cell(row=row,column=c)
        if cell.value is not None: cell.font=HEAD; cell.fill=PatternFill("solid",fgColor=NAVY)

def build_dre(ws):
    ws["A1"]="BISYOU — interactive channel DRE.  Edit the yellow cells; everything below recalculates (open in Excel)."
    ws["A1"].font=BOLD
    ws["A3"]="ASSUMPTIONS (edit yellow)"; ws["C3"]="Conservative"; ws["D3"]="Run-rate"; style_head(ws,3,4)
    A=[("Net revenue – annual (R$)",5_500_000,7_600_000,M),
       ("Gross margin %",0.80,0.80,P),
       ("Selling/logistics/card %",0.17,0.17,P),
       ("Marketing / CPA %",0.25,0.25,P),
       ("SG&A – annual (R$)",1_250_000,1_250_000,M),
       ("Royalty %",0.08,0.08,P),
       ("Tax gross-up % (net=gross*(1-tax))",0.13,0.13,P)]
    r=4
    for lab,cv,dv,fmt in A:
        ws.cell(row=r,column=1,value=lab)
        for col,val in ((3,cv),(4,dv)):
            cell=ws.cell(row=r,column=col,value=val); cell.number_format=fmt; cell.fill=INPUT_FILL; cell.font=BOLD
        r+=1
    # rows: 4 net,5 GM,6 sell,7 mkt,8 SGA,9 roy,10 tax
    ws["A12"]="DERIVED"; ws["A12"].font=BOLD
    for col in ("C","D"):
        ws[f"{col}12"]=f"={col}4/(1-{col}10)"; ws[f"{col}12"].number_format=M          # gross
        ws[f"{col}13"]=f"={col}12/12"; ws[f"{col}13"].number_format=M                    # monthly gross
        ws[f"{col}14"]=(f"=INDEX($G$31:$G$37,COUNT($F$31:$F$37)"
                        f"-COUNTIF($F$31:$F$37,\">\"&{col}13)+1)"); ws[f"{col}14"].number_format=P  # Boniteca %
    ws["A12"]="Gross revenue (R$)"; ws["A13"]="Monthly gross (R$)"; ws["A14"]="Boniteca % (regressive lookup)"
    # DRE
    ws["A16"]="DRE — ANNUAL RUN-RATE (R$)"; ws["C16"]="Conservative"; ws["D16"]="Run-rate"; style_head(ws,16,4)
    dre=[("Receita líquida","=C4","=D4",M,False),
         ("(–) CMV","=-(1-C5)*C4","=-(1-D5)*D4",M,False),
         ("= Margem bruta","=C5*C4","=D5*D4",M,True),
         ("(–) Despesas de venda","=-C6*C4","=-D6*D4",M,False),
         ("(–) Marketing / CPA","=-C7*C4","=-D7*D4",M,False),
         ("(–) SG&A","=-C8","=-D8",M,False),
         ("(–) Royalty","=-C9*C4","=-D9*D4",M,False),
         ("(–) Boniteca fee","=-C14*C12","=-D14*D12",M,False),
         ("= EBITDA (annual)","=C19+C20+C21+C22+C23+C24","=D19+D20+D21+D22+D23+D24",M,True),
         ("EBITDA margin %","=C25/C4","=D25/D4",P,False),
         ("= EBITDA — H2-2026 (Jul–Dec, ×0.5)","=C25*0.5","=D25*0.5",M,True)]
    r=17
    for lab,cf,df,fmt,bold in dre:
        ws.cell(row=r,column=1,value=lab)
        for col,f in ((3,cf),(4,df)):
            cell=ws.cell(row=r,column=col,value=f); cell.number_format=fmt
            if bold: cell.font=BOLD
        if bold: ws.cell(row=r,column=1).font=BOLD
        r+=1
    # Boniteca band table (mirrors Premissas r48-56)
    ws["F30"]="Boniteca band (monthly gross →)"; ws["G30"]="Fee %"; ws["F30"].font=BOLD; ws["G30"].font=BOLD
    bands=[(500000,0.10),(750000,0.09),(1000000,0.08),(1250000,0.075),(1500000,0.065),(2000000,0.06),(10000000,0.05)]
    rr=31
    for thr,rate in bands:
        ws.cell(row=rr,column=6,value=thr).number_format=M
        ws.cell(row=rr,column=7,value=rate).number_format=P
        rr+=1
    ws["A28"]="Boniteca uses GE's regressive table on MONTHLY gross (Bisyou's small scale → ~9%). Change Net revenue and the rate re-looks-up automatically."
    ws.column_dimensions['A'].width=40
    for c in ('B','C','D'): ws.column_dimensions[c].width=18
    ws.column_dimensions['F'].width=26; ws.column_dimensions['G'].width=10

def build_impact(ws,dre):
    q=f"'{dre}'!"
    ws["A1"]="GROUP P&L IMPACT — GE 2026 (BP) + Bisyou H2-2026 (live from the Bisyou tab, Conservative col)."; ws["A1"].font=BOLD
    ws["A3"]="Line"; ws["B3"]="GE 2026 (BP)"; ws["C3"]="Bisyou H2-2026"; ws["D3"]="Pro-forma"; ws["E3"]="Δ %"; style_head(ws,3,5)
    rows=[("Receita líquida",GE["Receita líquida"],f"={q}C17*0.5"),
          ("Margem bruta",GE["Margem bruta"],f"={q}C19*0.5"),
          ("Margem de contribuição",GE["Margem de contribuição"],f"=({q}C19+{q}C20+{q}C21)*0.5"),
          ("EBITDA",GE["EBITDA"],f"={q}C25*0.5")]
    r=4
    for lab,gev,bf in rows:
        ws.cell(row=r,column=1,value=lab)
        ws.cell(row=r,column=2,value=gev).number_format=M
        ws.cell(row=r,column=3,value=bf).number_format=M
        ws.cell(row=r,column=4,value=f"=B{r}+C{r}").number_format=M
        ws.cell(row=r,column=5,value=f"=C{r}/B{r}").number_format=P
        r+=1
    ws["A10"]="GE figures are the BP's 2026 group totals (static snapshot). Bisyou side is live from the Bisyou tab. NB: GE EBITDA is PRE-Boniteca (model convention); Bisyou EBITDA here is post-Boniteca+royalty."
    ws.column_dimensions['A'].width=26
    for c in ('B','C','D','E'): ws.column_dimensions[c].width=18

def build_sens(ws,dre):
    q=f"'{dre}'!"
    ws["A1"]="EBITDA CONTRIBUTION (annual, R$) — Royalty % (rows) × Net revenue (cols). LIVE: references the Bisyou tab's GM/selling/mktg/SG&A/tax + band table."; ws["A1"].font=BOLD
    revs=[5_000_000,5_500_000,6_000_000,6_500_000,7_000_000,7_600_000,8_000_000]
    roys=[0.03,0.04,0.05,0.06,0.07,0.08]
    ws["A3"]="Royalty / Net rev"; style_head(ws,3,1)
    for j,rv in enumerate(revs):
        c=ws.cell(row=3,column=2+j,value=rv); c.number_format=M; c.font=HEAD; c.fill=PatternFill("solid",fgColor=NAVY)
    for i,ro in enumerate(roys):
        ws.cell(row=4+i,column=1,value=ro).number_format='0%'; ws.cell(row=4+i,column=1).font=BOLD
        for j in range(len(revs)):
            colL=get_column_letter(2+j); rowN=4+i
            rev=f"{colL}$3"; roy=f"$A{rowN}"
            gross=f"({rev}/(1-{q}$C$10))"
            bon=(f"INDEX({q}$G$31:$G$37,COUNT({q}$F$31:$F$37)"
                 f"-COUNTIF({q}$F$31:$F$37,\">\"&({gross}/12))+1)")
            f=(f"={rev}*{q}$C$5-{rev}*{q}$C$6-{rev}*{q}$C$7-{q}$C$8-{roy}*{rev}-{bon}*{gross}")
            ws.cell(row=rowN,column=2+j,value=f).number_format=M
    # royalty ceiling row (live): r = GM - sell - mkt - SGA/rev - bon%(rev)
    ws.cell(row=12,column=1,value="Royalty ceiling (EBITDA=0)").font=BOLD
    for j in range(len(revs)):
        colL=get_column_letter(2+j); rev=f"{colL}$3"; gross=f"({rev}/(1-{q}$C$10))"
        bon=(f"INDEX({q}$G$31:$G$37,COUNT({q}$F$31:$F$37)-COUNTIF({q}$F$31:$F$37,\">\"&({gross}/12))+1)")
        f=(f"={q}$C$5-{q}$C$6-{q}$C$7-{q}$C$8/{rev}-{bon}*{gross}/{rev}")
        ws.cell(row=12,column=2+j,value=f).number_format=P
    # red/green conditional formatting on the grid
    rng=f"B4:{get_column_letter(1+len(revs))}9"
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan",formula=["0"],font=Font(color="B23B3B")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual",formula=["0"],font=Font(color="1F7A3F")))
    ws["A14"]="Green=accretive, red=dilutive (applied by Excel on open). Change marketing/SG&A on the Bisyou tab and the whole grid moves."
    ws.column_dimensions['A'].width=22
    for j in range(len(revs)): ws.column_dimensions[get_column_letter(2+j)].width=13

# ---------- STANDALONE ----------
wb=openpyxl.Workbook()
d=wb.active; d.title="Bisyou_DRE"; build_dre(d)
build_impact(wb.create_sheet("Group_Impact"),"Bisyou_DRE")
build_sens(wb.create_sheet("Sensitivity"),"Bisyou_DRE")
wb.save(STD); print("rebuilt standalone (interactive):",STD)

# ---------- BP COPY ----------
import shutil
shutil.copyfile(SRC,BPC)
bp=openpyxl.load_workbook(BPC)
for nm in ("Bisyou","Group_Impact","Sensitivity","Bisyou_DRE","BISYOU_Group_Impact","BISYOU_Sensitivity"):
    if nm in bp.sheetnames: del bp[nm]
build_dre(bp.create_sheet("Bisyou"))
build_impact(bp.create_sheet("Bisyou_GroupImpact"),"Bisyou")
build_sens(bp.create_sheet("Bisyou_Sensitivity"),"Bisyou")
mac=bp["Macro"]; mac["D4"]="Bisyou"; mac["B4"]=1
bp.save(BPC)
chk=openpyxl.load_workbook(BPC)
print("BP copy sheets:",len(chk.sheetnames),"| charts:",sum(len(getattr(chk[s],"_charts",[])) for s in chk.sheetnames),
      "| flags D1:D4:",[chk["Macro"][f"D{i}"].value for i in (1,2,3,4)])
