# -*- coding: utf-8 -*-
import openpyxl, shutil, re
SRC=r"c:\claude\sandbox\bisyou\diligence\Diligência Bisyou.xlsx"
OUT=r"c:\claude\sandbox\bisyou\diligence\Diligência Bisyou - PT.xlsx"
EM="—"; EN="–"; ARR="→"

T={
 "ASSUMPTIONS (edit yellow)":"PREMISSAS (edite as amarelas)",
 f"BISYOU {EM} interactive channel DRE.  Edit the yellow cells; everything below recalculates (open in Excel).":
   f"BISYOU {EM} DRE interativo do canal. Edite as células amarelas; tudo abaixo recalcula (abra no Excel).",
 "Bisyou % net":"Bisyou % líq.",
 "Boniteca % (regressive lookup)":"Boniteca % (busca regressiva)",
 f"Boniteca band (monthly gross {ARR})":f"Faixa Boniteca (bruto mensal {ARR})",
 f"Boniteca uses GE's regressive table on MONTHLY gross (Bisyou's small scale {ARR} ~9%). Change Net revenue and the rate re-looks-up automatically.":
   f"A Boniteca usa a tabela regressiva da GE sobre o bruto MENSAL (a escala menor da Bisyou {ARR} ~9%). Altere a Receita líquida e a taxa é rebuscada automaticamente.",
 "Conservative":"Conservador",
 f"DRE {EM} ANNUAL RUN-RATE (R$)":f"DRE {EM} RITMO ANUAL (R$)",
 "EBITDA margin %":"Margem EBITDA %",
 "Fee %":"Taxa %",
 "GE 2026 % (benchmark)":"GE 2026 % (referência)",
 f"GE BEAUTY {EM} macro P&L (quarterly) + savings from sharing Ops/Fin overhead with Bisyou":
   f"GE BEAUTY {EM} DRE macro (trimestral) + economia ao compartilhar a estrutura de Operações/Financeiro com a Bisyou",
 "GE P&L (as-is), R$":"DRE GE (atual), R$",
 f"GE P&L {EM} granular, per quarter (R$)":f"DRE GE {EM} detalhado, por trimestre (R$)",
 "Gross margin %":"Margem bruta %",
 f"Net revenue {EN} annual (R$)":f"Receita líquida {EN} anual (R$)",
 "Run-rate":"Ritmo atual",
 f"SG&A {EN} annual (R$)":f"SG&A {EN} anual (R$)",
 "Saving total":"Economia total",
 "Selling/logistics/card %":"Vendas/logística/cartão %",
 "Tax gross-up % (net=gross*(1-tax))":"Impostos gross-up % (líq.=bruto×(1-imposto))",
 "(=) EBITDA (annual)":"(=) EBITDA (anual)",
 f"({EN}) Boniteca fee":f"({EN}) Taxa Boniteca",
 "YTD":"Acum.",
 "Fee plataformas":"Taxa plataformas",
}
for y in (2026,2027):
    for q in (1,2,3,4): T[f"{y}Q{q}"]=f"{y} T{q}"
    T[f"{y} FY"]=f"{y} Anual"

shutil.copyfile(SRC,OUT)
wb=openpyxl.load_workbook(OUT,data_only=False)
n=0; per={}
for sh in wb.sheetnames:
    ws=wb[sh]
    for row in ws.iter_rows():
        for c in row:
            v=c.value
            if isinstance(v,str) and not v.startswith("=") and v in T:
                c.value=T[v]; n+=1; per[v]=per.get(v,0)+1
wb.save(OUT)
print("replacements:",n,"| distinct:",len(per))

# post-scan: any English leftovers?
EN_WORDS=re.compile(r"\b(annual|fee|Net revenue|Gross|Saving|Run-rate|benchmark|yellow|recalculat|lookup|Conservative|Selling|Tax gross|P&L|as-is|quarterly|overhead|monthly|interactive|channel|edit)\b",re.I)
chk=openpyxl.load_workbook(OUT,data_only=False)
left=set()
for sh in chk.sheetnames:
    for row in chk[sh].iter_rows():
        for c in row:
            if isinstance(c.value,str) and not c.value.startswith("=") and EN_WORDS.search(c.value):
                left.add((sh,c.coordinate,c.value))
print("\nEnglish leftovers:",len(left))
for x in sorted(left): print("  ",(str(x)).encode("ascii","replace").decode())
print("\nsheets:",len(chk.sheetnames),"| charts:",sum(len(getattr(chk[s],'_charts',[])) for s in chk.sheetnames))
# verify formulas intact (spot)
g=chk["GE_Macro_Savings"]
print("formula spot r18E:",str(g["E18"].value)[:60],"| r3 header C:",g["C3"].value,"| r38A:",g["A38"].value)
