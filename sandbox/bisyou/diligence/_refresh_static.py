# -*- coding: utf-8 -*-
import openpyxl, unicodedata
ORIG=r"G:\Drives compartilhados\GEB_Financeiro\Orçamento e Resultados\2026\GE Beauty_BP_v2026.xlsx"
COPY=r"C:\Users\Lucas Guimarães\Desktop\nami-works\sandbox\bisyou\diligence\GE Beauty_BP_v2026__+Bisyou.xlsx"
def norm(s): return unicodedata.normalize("NFKD",str(s)).encode("ascii","ignore").decode().strip().lower() if s is not None else ""

macv=openpyxl.load_workbook(ORIG,data_only=True)["Macro"]
mac_map={}
for r in range(1,macv.max_row+1):
    k=norm(macv.cell(row=r,column=4).value)
    if k: mac_map.setdefault(k,r)

# diligence col (C..L = 3..12) -> Macro month columns to sum
COL={3:[60,61,62],4:[63,64,65],5:[66,67,68],6:[69,70,71],7:list(range(60,72)),
     8:[79,80,81],9:[82,83,84],10:[85,86,87],11:[88,89,90],12:list(range(79,91))}

# (diligence_row -> macro line key). Disambiguated 'Sistemas': summary/adjusted = mkt-systems; granular SG&A = 'sistemas'
REFRESH={
 # block 1 (as-is summary)
 4:"receita bruta",5:"receita liquida",6:"cmv",8:"despesas de venda",9:"despesas de marketing",
 10:"sistemas de vendas e marketing",12:"sg&a",13:"financiamento estoque",15:"lucro liquido",
 # block 3 (adjusted) static cells only
 41:"receita bruta",42:"receita liquida",46:"despesas de marketing",47:"sistemas de vendas e marketing",52:"lucro liquido",
 # block 4 (granular source)
 68:"cmv",69:"margem bruta",70:"despesas variaveis",71:"despesas de venda",72:"armazenagem",
 73:"frete sobre vendas",74:"sistemas de vendas e marketing",75:"atendimento ao consumidor",
 76:"analytics",77:"mensalidade e-comm",78:"margem de contribuicao",79:"sg&a",80:"time clt",
 84:"sistemas",85:"erp",86:"e-mail + ferramentas",87:"escritorio",88:"servicos contratados",
 89:"ebitda",90:"financiamento estoque",92:"lucro liquido",
}
missing=[k for k in set(REFRESH.values()) if k not in mac_map]
print("Macro keys NOT found (skipped):",missing)

wbf=openpyxl.load_workbook(COPY,data_only=False)
g=wbf["GE_Macro_Savings"]
def mval(mr,cols):
    vs=[macv.cell(row=mr,column=c).value for c in cols]
    vs=[v for v in vs if isinstance(v,(int,float))]
    return sum(vs) if vs else None

updated=0; skipped_formula=0; sample=[]
for row,key in REFRESH.items():
    mr=mac_map.get(key)
    if mr is None: continue
    for c in range(3,13):
        cell=g.cell(row=row,column=c)
        cur=cell.value
        if isinstance(cur,str) and cur.startswith("="):
            skipped_formula+=1; continue          # never touch formulas
        if not isinstance(cur,(int,float)):
            continue                                # leave blanks untouched
        nv=mval(mr,COL[c])
        if nv is None: continue
        if len(sample)<14 and c in (6,7,12):
            sample.append((row,g.cell(row=row,column=1).value,c,round(cur),round(nv)))
        cell.value=nv; cell.number_format='#,##0'; updated+=1

# refresh Bisyou benchmark GE %s (static col F, rows 17-27)
b=wbf["Bisyou"]
nrev=mval(mac_map["receita liquida"],COL[7])
bench={17:"receita liquida",18:"cmv",19:"margem bruta",20:"despesas de venda",
       21:"despesas de marketing",22:"sg&a",24:"financiamento estoque",25:"ebitda"}
bcount=0
for rr,key in bench.items():
    cell=b.cell(row=rr,column=6)
    if isinstance(cell.value,str) and cell.value.startswith("="): continue
    val=1.0 if key=="receita liquida" else (mval(mac_map[key],COL[7])/nrev if key in mac_map else None)
    if val is not None: cell.value=val; cell.number_format='0.0%'; bcount+=1

wbf.save(COPY)
print("static GE cells updated:",updated,"| formulas skipped:",skipped_formula,"| benchmark %s refreshed:",bcount)
print("\nsample (row | line | col | OLD -> NEW), cols F=2026Q4,G=2026FY,L=2027FY:")
for row,lab,c,o,n in sample:
    print(("  r%d %-28s col%d  %s -> %s"%(row,str(lab),c,format(o,','),format(n,','))).encode("ascii","replace").decode())
