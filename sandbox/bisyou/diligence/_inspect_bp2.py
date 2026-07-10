# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.utils import get_column_letter

PATH = r"G:\Drives compartilhados\GEB_Financeiro\Orçamento e Resultados\2026\GE Beauty_BP_v2026.xlsx"
OUT = r"c:\claude\sandbox\bisyou\diligence\_bp_ecomm_skeleton.txt"

wv = openpyxl.load_workbook(PATH, data_only=True)["E-comm"]
wf = openpyxl.load_workbook(PATH, data_only=False)["E-comm"]

L = []
# 1) Full time header: rows 2 and 3 across ALL columns
for hr in (2, 3):
    parts = []
    for c in range(1, wv.max_column + 1):
        v = wv.cell(row=hr, column=c).value
        if v is not None:
            col = get_column_letter(c)
            if hasattr(v, "year") and hasattr(v, "month"):
                v = f"{v.year}-{v.month:02d}"
            parts.append(f"{col}={v}")
    L.append(f"--- ROW {hr} (time header) ---")
    L.append("  ".join(parts))

# 2) Label skeleton: rows 1..max, columns A-E text labels only
L.append("\n--- ROW LABELS (A | B | C | D | E) per row ---")
for r in range(1, wv.max_row + 1):
    a = wv.cell(row=r, column=1).value
    b = wv.cell(row=r, column=2).value
    c_ = wv.cell(row=r, column=3).value
    d = wv.cell(row=r, column=4).value
    e = wv.cell(row=r, column=5).value
    if any(x is not None for x in (a, b, c_, d, e)):
        def s(x): return "" if x is None else str(x)
        L.append(f"r{r}: A={s(a)} | B={s(b)} | C={s(c_)} | D={s(d)} | E={s(e)}")

open(OUT, "w", encoding="utf-8").write("\n".join(L))
print("wrote", OUT)
print("max_row", wv.max_row, "max_col", wv.max_column, "=", get_column_letter(wv.max_column))
