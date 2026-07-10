# -*- coding: utf-8 -*-
import openpyxl, unicodedata
ORIG=r"G:\Drives compartilhados\GEB_Financeiro\Orçamento e Resultados\2026\GE Beauty_BP_v2026.xlsx"
COPY=r"c:\claude\sandbox\bisyou\diligence\GE Beauty_BP_v2026__+Bisyou.xlsx"

def norm(s):
    if s is None: return ""
    return unicodedata.normalize("NFKD",str(s)).encode("ascii","ignore").decode().strip().lower()

macv=openpyxl.load_workbook(ORIG,data_only=True)["Macro"]   # pristine, cached values
wbf=openpyxl.load_workbook(COPY,data_only=False)            # write target (preserves formulas)
sav=wbf["GE_Macro_Savings"]

# label -> first Macro row (col D)
mac_map={}
for r in range(1,macv.max_row+1):
    d=macv.cell(row=r,column=4).value
    k=norm(d)
    if k: mac_map.setdefault(k,r)

ALIAS={"time":"time clt"}   # Lucas's short label -> Macro label
Q26={1:[60,61,62],2:[63,64,65],3:[66,67,68],4:[69,70,71]}
Q27={1:[79,80,81],2:[82,83,84],3:[85,86,87],4:[88,89,90]}
def qsum(r,cols):
    s=0.0; any_=False
    for c in cols:
        v=macv.cell(row=r,column=c).value
        if isinstance(v,(int,float)): s+=v; any_=True
    return s if any_ else None

# header for the granular block (row 17, currently empty)
hdr=["GE P&L — granular, per quarter (R$)","2026Q1","2026Q2","2026Q3","2026Q4","2026 FY","2027Q1","2027Q2","2027Q3","2027Q4","2027 FY"]
for j,h in enumerate(hdr):
    c=sav.cell(row=17,column=1+j,value=h)
    from openpyxl.styles import Font,PatternFill
    c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="11243F")

matched=[]; left=[]
for lr in range(18,50):
    lab=sav.cell(row=lr,column=1).value
    if not lab: continue
    if "%" in str(lab):
        left.append((lr,lab,"ratio — skipped")); continue
    k=norm(lab); k=ALIAS.get(k,k)
    mr=mac_map.get(k)
    if mr is None:
        left.append((lr,lab,"not in GE data — left blank")); continue
    q26=[qsum(mr,Q26[i]) for i in (1,2,3,4)]
    q27=[qsum(mr,Q27[i]) for i in (1,2,3,4)]
    fy26=sum(v for v in q26 if v is not None) if any(v is not None for v in q26) else None
    fy27=sum(v for v in q27 if v is not None) if any(v is not None for v in q27) else None
    vals=q26+[fy26]+q27+[fy27]
    for j,v in enumerate(vals):
        if v is not None:
            sav.cell(row=lr,column=2+j,value=v).number_format='#,##0'
    matched.append((lr,lab,mr))

wbf.save(COPY)
print("POPULATED (%d):"%len(matched))
for lr,lab,mr in matched: print(("  r%d  %-32s <- Macro r%d"%(lr,lab,mr)).encode("ascii","replace").decode())
print("\nLEFT AS-IS (%d):"%len(left))
for lr,lab,why in left: print(("  r%d  %-32s (%s)"%(lr,lab,why)).encode("ascii","replace").decode())
