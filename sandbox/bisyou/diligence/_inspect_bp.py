# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.utils import get_column_letter

PATH = r"G:\Drives compartilhados\GEB_Financeiro\Orçamento e Resultados\2026\GE Beauty_BP_v2026.xlsx"
OUT = r"C:\Users\Lucas Guimarães\Desktop\nami-works\sandbox\bisyou\diligence\_bp_ecomm_dump.txt"

wb_v = openpyxl.load_workbook(PATH, data_only=True)
wb_f = openpyxl.load_workbook(PATH, data_only=False)

lines = []
lines.append("SHEETS: " + " | ".join(wb_v.sheetnames))

# find the e-comm sheet
target = None
for name in wb_v.sheetnames:
    n = name.lower().replace("-", "").replace(" ", "")
    if "ecomm" in n or "ecommerce" in n:
        target = name; break
lines.append("TARGET SHEET: " + str(target))

if target:
    wv = wb_v[target]; wf = wb_f[target]
    lines.append(f"DIMS: {wv.dimensions} | max_row={wv.max_row} max_col={wv.max_column}")
    maxc = min(wv.max_column, 45)
    maxr = min(wv.max_row, 400)
    for r in range(1, maxr + 1):
        cells = []
        for c in range(1, maxc + 1):
            v = wv.cell(row=r, column=c).value
            f = wf.cell(row=r, column=c).value
            if v is None and f is None:
                continue
            col = get_column_letter(c)
            # show formula if the cell is a formula (starts with =)
            if isinstance(f, str) and f.startswith("="):
                val = "" if v is None else (round(v, 2) if isinstance(v, float) else v)
                cells.append(f"{col}{r}[{f} ->{val}]")
            else:
                val = round(v, 4) if isinstance(v, float) else v
                cells.append(f"{col}{r}={val}")
        if cells:
            lines.append("  ".join(cells))

open(OUT, "w", encoding="utf-8").write("\n".join(lines))
print("wrote", OUT, "rows dumped:", len(lines))
print("sheets:", wb_v.sheetnames)
print("target:", target)
